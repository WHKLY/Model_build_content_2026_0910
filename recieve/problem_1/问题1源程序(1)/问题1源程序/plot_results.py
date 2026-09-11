#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""只读取问题 1 的结果 xlsx，绘制购电量和储电量，不重新求解模型。

安装依赖：
    python -m pip install numpy matplotlib openpyxl

用法（在 VS Code 终端中运行）：
    python plot_results.py --file "result1.xlsx" --output-dir "figures"

默认读取 solve_q1.py 写出的“逐时明细”工作表。首行为列名，必须有：
    时段、计划购电量、期初储电量、期末储电量
每一行代表一个 10 分钟时段，全天恰有 144 行，时段覆盖 00:00—24:00。
购电量单位为 kWh/10 分钟，储电量单位为 kWh。

原始空白模板只有每 4 小时的充放电汇总，不能据此唯一还原 10 分钟的储电量；
因此本程序不会对模板中的汇总数据擅自插值或重新优化。
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path
import re
import warnings

import numpy as np
from openpyxl import load_workbook


REQUIRED_COLUMNS = ("时段", "计划购电量", "期初储电量", "期末储电量")
SOC_MIN = 1200.0
SOC_MAX = 10800.0
TOL = 1e-5


def minutes_of_day(value: str) -> int:
    """将 23:50、24:00、0:00+1 等明确时间标签转为距当天零点的分钟数。"""
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::00)?(?:\+(\d+))?\s*", value)
    if not match:
        raise ValueError(f"无法识别时间标签：{value!r}")
    hour, minute, day = (int(part or 0) for part in match.groups())
    if not (0 <= hour <= 24 and 0 <= minute < 60) or (hour == 24 and minute != 0):
        raise ValueError(f"时间标签越界：{value!r}")
    return day * 1440 + hour * 60 + minute


def parse_period(value: object) -> tuple[int, int]:
    text = str(value).strip().replace("—", "-").replace("–", "-").replace("～", "-")
    parts = text.split("-")
    if len(parts) != 2:
        raise ValueError(f"时段须为 0:00-0:10 这样的格式，实际为：{value!r}")
    return minutes_of_day(parts[0]), minutes_of_day(parts[1])


def as_number(value: object, cell_label: str) -> float:
    if value is None or isinstance(value, bool):
        raise ValueError(f"{cell_label} 缺少有效数值。请读取求解程序写出的结果文件。")
    try:
        result = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{cell_label} 不是数值：{value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{cell_label} 包含 NaN 或无穷大。")
    return result


def read_result(path: Path, sheet_name: str) -> tuple[np.ndarray, np.ndarray]:
    """读取并核对时段顺序和 SOC 连续性；不修改任何 xlsx。"""
    if not path.is_file():
        raise ValueError(f"结果文件不存在：{path}")
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in book.sheetnames:
            raise ValueError(
                f"文件缺少“{sheet_name}”工作表。请先运行问题 1 求解代码。"
                "原模板的 4 小时充放电汇总不能唯一确定 10 分钟储电量曲线。"
            )
        sheet = book[sheet_name]
        row_iter = sheet.iter_rows(values_only=True)
        first = next(row_iter, None)
        if first is None:
            raise ValueError(f"“{sheet_name}”工作表为空。")
        header = [str(value).strip() if value is not None else "" for value in first]
        for name in REQUIRED_COLUMNS:
            if header.count(name) != 1:
                raise ValueError(f"首行必须且只能包含一列“{name}”，实际列名为：{header}")
        columns = {name: header.index(name) for name in REQUIRED_COLUMNS}
        records = []
        for excel_row, values in enumerate(row_iter, start=2):
            if all(value is None for value in values):
                continue
            period = parse_period(values[columns["时段"]])
            numbers = [
                as_number(values[columns[name]], f"{sheet_name} 第 {excel_row} 行“{name}”")
                for name in REQUIRED_COLUMNS[1:]
            ]
            records.append((period, *numbers))
    finally:
        book.close()

    if len(records) != 144:
        raise ValueError(f"问题 1 应有 144 个时段，实际有 {len(records)} 行；请检查工作表。")
    records.sort(key=lambda row: row[0][0])
    for index, record in enumerate(records):
        expected = (index * 10, (index + 1) * 10)
        if record[0] != expected:
            raise ValueError(
                f"第 {index + 1} 个时段应覆盖距零点 {expected[0]}—{expected[1]} 分钟，"
                f"实际为 {record[0]}。请检查缺失、重复、跨日或整体错移的标签。"
            )
    data = np.asarray([row[1:] for row in records], dtype=float)
    grid, start_soc, end_soc = data.T
    if np.min(grid) < -TOL:
        raise ValueError("购电量出现负值；本模型仅计非负购电量，请先核查结果。")
    if not np.allclose(start_soc[1:], end_soc[:-1], atol=TOL, rtol=0):
        raise ValueError("相邻时段的期末与期初储电量不一致，无法绘制连续 SOC 曲线。")
    soc = np.concatenate(([start_soc[0]], end_soc))  # 144 个区间对应 145 个边界点。
    if soc.min() < SOC_MIN - TOL or soc.max() > SOC_MAX + TOL:
        raise ValueError("储电量超出题设的 1200—10800 kWh 范围，请先检查求解结果。")
    if abs(soc[-1] - soc[0]) > TOL:
        raise ValueError("0:00 与 24:00 的储电量不同，不满足问题 1 的日循环约束。")
    return grid, soc


def configure_font(matplotlib) -> bool:
    """优先选择已安装的中文字体；没有时自动改用英文图签，避免中文方框。"""
    from matplotlib import font_manager

    candidates = (
        "Microsoft YaHei", "Microsoft YaHei UI", "SimHei", "Noto Sans CJK SC",
        "Source Han Sans SC", "PingFang SC", "WenQuanYi Micro Hei", "SimSun",
        "Arial Unicode MS",
    )
    installed = {font.name for font in font_manager.fontManager.ttflist}
    for name in candidates:
        if name in installed:
            matplotlib.rcParams["font.family"] = [name, "DejaVu Sans"]
            matplotlib.rcParams["axes.unicode_minus"] = False
            return True
    matplotlib.rcParams["font.family"] = "DejaVu Sans"
    warnings.warn("未找到常用中文字体，图片将自动使用英文图签。", stacklevel=2)
    return False


def create_plot(grid: np.ndarray, soc: np.ndarray, output_path: Path, dpi: int) -> None:
    import matplotlib

    matplotlib.use("Agg")  # 在 VS Code 终端和无显示器环境中均可保存图片。
    from matplotlib import pyplot as plt
    from matplotlib.ticker import FuncFormatter

    chinese = configure_font(matplotlib)
    text = {
        "title": "问题 1：计划购电量与储电量" if chinese else "Problem 1: Grid purchase and stored energy",
        "grid": "计划购电量" if chinese else "Planned grid purchase",
        "soc": "储电量" if chinese else "Stored energy",
        "grid_ylabel": "购电量（kWh / 10分钟）" if chinese else "Grid purchase (kWh / 10 min)",
        "soc_ylabel": "储电量（kWh）" if chinese else "Stored energy (kWh)",
        "xlabel": "时刻" if chinese else "Time of day",
        "lower": "下限 1,200 kWh" if chinese else "Lower bound: 1,200 kWh",
        "upper": "上限 10,800 kWh" if chinese else "Upper bound: 10,800 kWh",
    }
    plt.rcParams.update({"font.size": 11, "axes.titlesize": 15, "axes.labelsize": 11})
    fig, (ax_grid, ax_soc) = plt.subplots(
        2, 1, sharex=True, figsize=(12, 7.2), layout="constrained",
        gridspec_kw={"height_ratios": [1, 1.15]},
    )
    hours = np.arange(145, dtype=float) / 6.0
    # 每个阶梯覆盖完整 10 分钟区间，避免把区间电量画在错误的时间位置。
    ax_grid.stairs(grid, hours, fill=True, color="#26729C", alpha=0.18)
    ax_grid.stairs(grid, hours, color="#18638B", linewidth=1.5, label=text["grid"])
    ax_grid.set_title(text["title"], loc="left", pad=13)
    ax_grid.set_ylabel(text["grid_ylabel"])
    ax_grid.set_ylim(0, max(1.0, float(grid.max()) * 1.12))
    ax_grid.legend(loc="upper right", frameon=False)

    # 恒功率离散模型下，SOC 在每个 10 分钟区间内线性变化，因此连接边界点。
    ax_soc.axhspan(SOC_MIN, SOC_MAX, color="#E6F1EC", alpha=0.75, zorder=0)
    ax_soc.plot(hours, soc, color="#26724B", linewidth=2.0, label=text["soc"])
    ax_soc.scatter([0, 24], [soc[0], soc[-1]], color="#26724B", s=30, zorder=5)
    ax_soc.axhline(SOC_MIN, color="#9C6651", linestyle="--", linewidth=1.1, label=text["lower"])
    ax_soc.axhline(SOC_MAX, color="#8B5969", linestyle="--", linewidth=1.1, label=text["upper"])
    ax_soc.set_ylabel(text["soc_ylabel"])
    ax_soc.set_xlabel(text["xlabel"])
    ax_soc.set_ylim(0, 12000)
    ax_soc.set_yticks([0, SOC_MIN, 3000, 6000, 9000, SOC_MAX, 12000])
    ax_soc.legend(loc="upper center", bbox_to_anchor=(0.5, -0.24), ncols=3, frameon=False)
    ax_soc.set_xlim(0, 24)
    ax_soc.set_xticks(np.arange(0, 25, 2))
    ax_soc.xaxis.set_major_formatter(FuncFormatter(lambda hour, _: f"{int(hour):02d}:00"))
    for axis in (ax_grid, ax_soc):
        axis.grid(axis="y", color="#CBD5E1", linewidth=0.6, alpha=0.75)
        axis.set_axisbelow(True)
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(length=3, color="#94A3B8")
        axis.yaxis.set_major_formatter(FuncFormatter(lambda value, _: f"{value:,.0f}"))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        fig.savefig(output_path, dpi=dpi, bbox_inches="tight", facecolor="white")
    finally:
        plt.close(fig)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", type=Path, required=True, help="求解代码生成的 result1.xlsx 路径")
    parser.add_argument("--sheet", default="逐时明细", help="10 分钟明细表名称，默认：逐时明细")
    parser.add_argument("--output-dir", type=Path, default=Path("figures"), help="图片保存目录")
    parser.add_argument("--dpi", type=int, default=300, help="图片分辨率，默认 300 dpi")
    args = parser.parse_args()
    if not 72 <= args.dpi <= 1200:
        parser.error("--dpi 应在 72 至 1200 之间。")
    try:
        grid, soc = read_result(args.file, args.sheet)
        output = args.output_dir / f"{args.file.stem}_purchase_soc.png"
        create_plot(grid, soc, output, args.dpi)
    except (ValueError, OSError) as exc:
        parser.exit(2, f"绘图失败：{exc}\n")
    print(f"图片已保存：{output.resolve()}")


if __name__ == "__main__":
    main()
