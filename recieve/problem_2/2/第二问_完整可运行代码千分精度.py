# -*- coding: utf-8 -*-
"""第二问 · 整合单文件可运行脚本（由《第二问代码审阅稿》整合修订）。

物理口径：计划电允许未取用但全额付费；只允许富余电充电；当段功率近似恒定并实时可测；
未来仅使用零点因果预测；紧急购电按 5 倍电价计费且只能补负荷、不能给电池充电；
终端 SOC 软惩罚只进决策、不进真实费用统计。

相对审阅稿的整合修订：
1) 附件位置大修正：删除写死的 "F:/…/C题/附件"，改为以本脚本所在目录为默认输入目录，
   并依次回退 输入目录/附件、旧路径；result2.xlsx 官方空模板只读不覆盖，结果另存输出目录。
2) 单文件化：去掉 from q2_forecast import …、display()、%pip 等 Notebook/跨模块写法。
3) 附件1 只读第0、1列（时间、电价），忽略第2、3列（小区负载、光伏发电预测功率，属其他问）；
   时间列兼容 小数天/“h:mm”文本/“0:00+1” 三种存储形式。
4) 一月预热可选：JAN_WARMUP=True 时从 1 月 1 日连续仿真，SOC 自然传递到 2 月 1 日。
5) 预留可选 DeepSeek API 接口（仅生成结果分析文字），核心计算完全本地、无需联网与 Key。

依赖：numpy / pandas / scipy / openpyxl（pip install numpy pandas scipy openpyxl）。
"""
from __future__ import annotations

import json
import urllib.request
import warnings
from dataclasses import dataclass, asdict
from datetime import datetime, time, timedelta
from pathlib import Path
from time import perf_counter

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from scipy.optimize import linprog, milp, Bounds, LinearConstraint
from scipy.sparse import coo_matrix, csr_matrix, diags, hstack, vstack

# ===================== 0. 用户配置区（只需要改这里） =====================
SCRIPT_DIR = Path(__file__).resolve().parent          # 本脚本所在目录
INPUT_DIR = SCRIPT_DIR                                # ← 改成附件1.xlsx/附件2.xlsx所在目录
OUTPUT_DIR = SCRIPT_DIR / "第二问输出"                 # ← 结果输出目录（自动创建）
LEGACY_INPUT_DIR = Path("F:/数学建模/26.9数模国赛2.0/C题/附件")  # 仅最后回退，可留可删
RESULT_TEMPLATE = INPUT_DIR / "result2.xlsx"          # 官方空模板：只读表头，绝不覆盖
USE_TEMPLATE_HEADER = False   # True=逐字复制官方模板147列表头；False=规范区间表头+时间映射表
JAN_WARMUP = True             # True=1月1日起连续仿真、SOC传递到2月1日；False=2月1日重置6000
OUTPUT_START = "2025-02-01"   # 正式输出起始日（1月为预热/历史期）
# ---- 可选：DeepSeek API（仅用于生成结果分析文字；留空则自动跳过） ----
USE_DEEPSEEK_REPORT = False
DEEPSEEK_API_KEY = ""         # ← 如需要，填你的 Key，形如 "sk-xxxxxxxx"
DEEPSEEK_BASE_URL = "https://api.deepseek.com"
DEEPSEEK_MODEL = "deepseek-chat"

SLOTS_PER_DAY = 144           # [D1] T=24×60/10
DT_HOURS = 1.0 / 6.0          # [D1] Δt=10/60 小时
SPECIAL_DATES = ("2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21")
EPS = 1e-7


# ===================== 1. [D1] 数据校验与附件读取 =====================
def _time_to_minute(value) -> int:
    """[D1] 将 00:10…24:00（含小数天、文本、次日零点）统一成 10…1440 分钟。"""
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        assert value.second == 0 and value.microsecond == 0, "时间标签必须落在整分钟。"
        return value.hour * 60 + value.minute
    if isinstance(value, timedelta):
        minutes = value.total_seconds() / 60.0
        assert minutes == int(minutes), "时间标签必须落在整分钟。"
        return int(minutes)
    if isinstance(value, (int, float, np.number)):      # 兼容 Excel 小数天存储
        minutes = float(value) * 1440.0
        assert abs(minutes - round(minutes)) < 1e-5, "无法解析小数时间标签。"
        return int(round(minutes))
    label = str(value).strip().replace("：", ":").replace(" ", "")
    next_day = "+1" in label or "次日" in label
    label = label.replace("+1", "").replace("次日", "")
    parts = label.split(":")
    assert len(parts) in (2, 3), f"无法解析时间标签：{value!r}"
    hour, minute = int(parts[0]), int(parts[1])
    assert len(parts) == 2 or int(parts[2]) == 0, "时间标签秒数必须为0。"
    assert 0 <= hour <= 24 and 0 <= minute < 60, f"无效时间：{value!r}"
    assert hour < 24 or minute == 0, "24时只能为24:00。"
    return hour * 60 + minute + (1440 if next_day else 0)


def _validate_time_labels(labels, name: str) -> np.ndarray:
    """[D1] 标签必须严格按 00:10…24:00 排列。"""
    actual = np.asarray([_time_to_minute(x) for x in labels], dtype=int)
    expected = np.arange(10, 1441, 10, dtype=int)
    assert actual.shape == expected.shape, f"{name}应有144个时段，实际为{actual.size}。"
    assert np.array_equal(actual, expected), f"{name}时间顺序须为00:10、00:20、…、24:00。"
    return actual


def _positive_spike_count(values: np.ndarray) -> int:
    """[D1] 孤立正尖峰审计，只计数不修改数据。"""
    center = values[:, 1:-1]
    neighbor = np.maximum(values[:, :-2], values[:, 2:])
    mask = (center > 100.0) & (center > 10.0 * np.maximum(neighbor, 1.0))
    return int(mask.sum())


def validate_data(load, pv, price, dates=None) -> pd.DataFrame:
    """[D1] 形状、有限性、非负、正电价与日期连续性校验，返回审计表。"""
    load = np.asarray(load, dtype=float)
    pv = np.asarray(pv, dtype=float)
    price = np.asarray(price, dtype=float)
    assert load.ndim == 2 and load.shape[1] == SLOTS_PER_DAY and load.shape[0] > 0, "负荷必须为非空的(日数,144)数组。"
    assert pv.shape == load.shape, "光伏与负荷数组形状必须完全一致。"
    assert price.shape == (SLOTS_PER_DAY,), "电价必须恰为144段。"
    for name, arr in (("load", load), ("pv", pv), ("price", price)):
        assert np.isfinite(arr).all(), f"{name}含NaN或无穷值，已拒绝运行。"
        assert (arr >= 0).all(), f"{name}含负值，已拒绝运行。"
    assert (price > 0).all(), "电价必须严格为正。"
    if dates is not None:
        idx = pd.DatetimeIndex(pd.to_datetime(dates, errors="raise"))
        assert len(idx) == load.shape[0], "日期数与数据天数不一致。"
        assert not idx.hasnans and idx.is_unique, "日期含缺失或重复。"
        assert idx.equals(idx.normalize()), "日期列须为每天零点。"
        assert idx.is_monotonic_increasing, "日期必须严格递增。"
        assert np.all((idx[1:] - idx[:-1]) == pd.Timedelta(days=1)), "日期必须每日连续。"
    records = []
    for name, arr in (("load", load), ("pv", pv), ("price", price)):
        records.append({"item": name, "shape": str(arr.shape), "min": float(arr.min()),
                        "max": float(arr.max()), "missing": int(np.isnan(arr).sum()),
                        "negative": int((arr < 0).sum()),
                        "isolated_positive_spikes": _positive_spike_count(arr) if arr.ndim == 2 else 0,
                        "status": "PASS"})
    records.append({"item": "pv_gt_load_fraction", "shape": str(load.shape),
                    "min": float(np.mean(pv > load)), "max": float(np.mean(pv > load)),
                    "missing": 0, "negative": 0, "isolated_positive_spikes": 0, "status": "INFO"})
    return pd.DataFrame(records)


def _resolve_input_file(input_dir: Path, filename: str) -> Path:
    """[D1] 附件定位：输入目录→输入目录/附件→脚本目录→脚本目录/附件→旧路径。"""
    candidates = [input_dir / filename, input_dir / "附件" / filename,
                  SCRIPT_DIR / filename, SCRIPT_DIR / "附件" / filename,
                  LEGACY_INPUT_DIR / filename, LEGACY_INPUT_DIR / "附件" / filename]
    for candidate in candidates:
        if candidate.is_file():
            return candidate
    raise FileNotFoundError(f"未找到{filename}；已搜索：{[str(c) for c in candidates]}")


def _read_power_sheet(sheet) -> tuple:
    """[D1] 附件2的 日期×功率(kW) 矩阵 → 日期×电量(kWh) 矩阵。"""
    rows = list(sheet.iter_rows(values_only=True))
    assert len(rows) >= 2 and len(rows[0]) == 145, f"{sheet.title}应有1个日期列和144个功率列。"
    labels = list(rows[0][1:])
    _validate_time_labels(labels, sheet.title)
    dates = pd.DatetimeIndex(pd.to_datetime([row[0] for row in rows[1:]], errors="raise"))
    power = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    return power * DT_HOURS, dates, labels          # ℓ=P_load/6，v=P_PV/6


def load_inputs(input_dir: Path = None) -> dict:
    """[D1] 读取附件1（电价）与附件2（负荷/光伏），返回数组、审计与时间映射。只读不改原件。"""
    input_dir = Path(INPUT_DIR if input_dir is None else input_dir).expanduser().resolve()
    price_path = _resolve_input_file(input_dir, "附件1.xlsx")
    data_path = _resolve_input_file(input_dir, "附件2.xlsx")
    wb_price = load_workbook(price_path, data_only=True, read_only=True)
    try:
        price_rows = list(wb_price.worksheets[0].iter_rows(values_only=True))
        assert len(price_rows) == 145, "附件1必须有表头与144个电价时段。"
        _validate_time_labels([row[0] for row in price_rows[1:]], "附件1")
        price = np.asarray([row[1] for row in price_rows[1:]], dtype=float)  # 只取第1列电价
    finally:
        wb_price.close()
    wb_data = load_workbook(data_path, data_only=True, read_only=True)
    try:
        assert len(wb_data.worksheets) >= 2, "附件2缺少负荷或光伏工作表。"
        load_sheet = next((s for s in wb_data.worksheets if "负荷" in s.title or "负载" in s.title), None)
        pv_sheet = next((s for s in wb_data.worksheets if "光伏" in s.title), None)
        assert load_sheet is not None and pv_sheet is not None, "附件2必须包含负荷表及光伏表。"
        load, dates, labels = _read_power_sheet(load_sheet)
        pv, pv_dates, pv_labels = _read_power_sheet(pv_sheet)
        assert dates.equals(pv_dates), "负荷与光伏的日期行没有一一对齐。"
        assert [_time_to_minute(x) for x in labels] == [_time_to_minute(x) for x in pv_labels], "负荷和光伏时段标签不同。"
    finally:
        wb_data.close()
    audit = validate_data(load, pv, price, dates)
    assert dates.equals(pd.date_range("2025-01-01", "2025-12-31", freq="D")), "正式输入应覆盖2025年完整365天。"
    minute_end = np.arange(10, 1441, 10)
    time_map = pd.DataFrame({
        "slot_index_0based": np.arange(144), "slot_t_1based": np.arange(1, 145),
        "source_endpoint_minute": minute_end,
        "source_endpoint": [f"{m // 60:02d}:{m % 60:02d}" for m in minute_end],
        "physical_interval": [f"{(m - 10) // 60:02d}:{(m - 10) % 60:02d}-{m // 60:02d}:{m % 60:02d}" for m in minute_end],
        "power_to_energy_factor": DT_HOURS})
    print(f"[D1] 读取{load.shape[0]}天×{load.shape[1]}段；电量kWh、电价元/kWh；光伏>负荷占比{np.mean(pv > load):.4%}。")
    return {"load": load, "pv": pv, "price": price, "dates": dates,
            "input_audit": audit, "time_map": time_map,
            "source_paths": {"price": str(price_path), "actual": str(data_path)}}


# ===================== 2. [F1-F3] 因果预测与联合误差情景 =====================
def prepare_forecasts(load, pv, window: int = 7) -> dict:
    """[F1][F2] 逐日因果均值预测与日结束后残差；第 d 天只读 d 之前的数据。"""
    load = np.asarray(load, dtype=float)
    pv = np.asarray(pv, dtype=float)
    validate_data(load, pv, np.ones(SLOTS_PER_DAY))
    assert isinstance(window, (int, np.integer)) and window >= 1, "window必须为正整数。"
    load_hat = np.zeros_like(load)
    pv_hat = np.zeros_like(pv)
    history_start = np.zeros(load.shape[0], dtype=int)
    for day in range(1, load.shape[0]):
        start = max(0, day - window)
        load_hat[day] = load[start:day].mean(axis=0)
        pv_hat[day] = pv[start:day].mean(axis=0)
        history_start[day] = start
    load_err = load - load_hat
    pv_err = pv - pv_hat
    assert np.isfinite(load_hat).all() and np.isfinite(pv_hat).all(), "预测含非有限值。"
    print(f"[F1] 已完成{load.shape[0]}天预测；窗口={window}天；第d天仅用[:d]；第0天预测为0。")
    return {"load_hat": load_hat, "pv_hat": pv_hat, "load_err": load_err, "pv_err": pv_err,
            "window": int(window), "history_start": history_start}


def make_scenarios(day, forecast, load, pv, scenario_count=20, history_days=28,
                   seed=20260911, residual_scale=1.0):
    """[F3] 不放回抽取过去完整日联合误差；空池时返回点预测回退情景(-1标记)。"""
    assert isinstance(day, (int, np.integer)) and 0 <= day < len(load), "day超出数据范围。"
    assert np.shape(load) == np.shape(pv) == np.shape(forecast["load_hat"]) == np.shape(forecast["pv_hat"]), "预测与实际形状不一致。"
    assert scenario_count > 0 and history_days > 0 and residual_scale >= 0, "情景参数非法。"
    start = max(int(forecast.get("window", 7)), day - int(history_days))
    pool = np.arange(start, day, dtype=int)
    if pool.size == 0:
        source_days = np.array([-1], dtype=int)
        scenario_load = np.asarray(forecast["load_hat"][day], dtype=float)[None, :].copy()
        scenario_pv = np.asarray(forecast["pv_hat"][day], dtype=float)[None, :].copy()
    else:
        rng = np.random.default_rng(int(seed) + int(day))
        count = min(int(scenario_count), int(pool.size))
        source_days = np.sort(rng.choice(pool, size=count, replace=False))
        scenario_load = np.maximum(0.0, forecast["load_hat"][day][None, :] + residual_scale * forecast["load_err"][source_days])
        scenario_pv = np.maximum(0.0, forecast["pv_hat"][day][None, :] + residual_scale * forecast["pv_err"][source_days])
    weights = np.full(len(source_days), 1.0 / len(source_days))
    assert np.all(source_days < day), "情景池发生未来信息泄漏。"
    assert scenario_load.shape == scenario_pv.shape == (len(weights), SLOTS_PER_DAY), "联合情景维度错误。"
    assert np.isfinite(scenario_load).all() and np.isfinite(scenario_pv).all(), "情景含非有限值。"
    assert (scenario_load >= 0).all() and (scenario_pv >= 0).all() and np.isclose(weights.sum(), 1.0), "情景取值或概率不合法。"
    return scenario_load, scenario_pv, weights, source_days


def point_summary(forecast, load, pv, dates=None, start_day=31) -> pd.DataFrame:
    """[F2] 逐日预测误差指标（MAE/RMSE），统计不参与参数拟合。"""
    count = len(load)
    assert 0 <= start_day < count, "start_day超出评价范围。"
    indices = pd.DatetimeIndex(dates) if dates is not None else np.arange(count)
    el = np.asarray(load)[start_day:] - forecast["load_hat"][start_day:]
    ev = np.asarray(pv)[start_day:] - forecast["pv_hat"][start_day:]
    return pd.DataFrame({"date": indices[start_day:],
                         "load_MAE_kWh": np.mean(np.abs(el), axis=1),
                         "load_RMSE_kWh": np.sqrt(np.mean(el ** 2, axis=1)),
                         "pv_MAE_kWh": np.mean(np.abs(ev), axis=1),
                         "pv_RMSE_kWh": np.sqrt(np.mean(ev ** 2, axis=1)),
                         "net_MAE_kWh": np.mean(np.abs(el - ev), axis=1),
                         "net_RMSE_kWh": np.sqrt(np.mean((el - ev) ** 2, axis=1))})


def scenario_summary(day, scenario_load, scenario_pv, weights, source_days, dates=None) -> pd.DataFrame:
    """[F3] 情景来源/概率/日电量审计表。"""
    source_days = np.asarray(source_days, dtype=int)
    assert np.all(source_days < day), "摘要检查发现情景来源不是过去日期。"
    labels = ["point_fallback" if i < 0 else (str(pd.Timestamp(dates[i]).date()) if dates is not None else str(i))
              for i in source_days]
    return pd.DataFrame({"target_day": day, "scenario": np.arange(1, len(weights) + 1),
                         "source_day": source_days, "source_date": labels, "probability": weights,
                         "load_total_kWh": np.asarray(scenario_load).sum(axis=1),
                         "pv_total_kWh": np.asarray(scenario_pv).sum(axis=1),
                         "net_total_kWh": (np.asarray(scenario_load) - np.asarray(scenario_pv)).sum(axis=1)})


# ===================== 3. [M/S] 日前两阶段随机优化 =====================
@dataclass
class Config:
    dt: float = 1 / 6
    eta_c: float = 0.9
    eta_d: float = 0.9
    s_min: float = 1200.0
    s_max: float = 10800.0
    p_max: float = 5000.0
    s_initial: float = 6000.0
    s_ref: float = 6000.0
    kappa: float = 0.5
    emergency_multiplier: float = 5.0
    window: int = 7
    history_days: int = 28
    scenarios: int = 20
    seed: int = 20260911
    mip_time_limit: float = 30.0
    mip_rel_gap: float = 1e-3
    energy_tol: float = 2e-5
    residual_scale: float = 1.0

    def validate(self):
        assert 0 < self.eta_c <= 1 and 0 < self.eta_d <= 1, "充放电效率必须在(0,1]内"
        assert 0 <= self.s_min < self.s_max, "储电量区间无效"
        assert self.s_min <= self.s_initial <= self.s_max, "初始SOC超出设备范围"
        assert self.s_min <= self.s_ref <= self.s_max, "参考SOC超出设备范围"
        assert self.p_max > 0 and self.dt > 0, "功率或步长无效"
        assert self.kappa >= 0 and self.emergency_multiplier > 1, "费用参数无效"
        assert self.window >= 1 and self.history_days >= 1 and self.scenarios >= 1, "历史窗口无效"
        assert np.isfinite(np.array(list(asdict(self).values()), dtype=float)).all(), "参数非有限值"
        return self


def _check_arrays(load, pv, price, s0, cfg):
    cfg.validate()
    for name, value in [("负载", load), ("光伏", pv), ("电价", price)]:
        assert np.isfinite(value).all(), f"{name}包含NaN或Inf"
        assert np.min(value) >= 0, f"{name}含负值"
    assert np.all(np.asarray(price) > 0), "该模型要求正电价"
    assert np.shape(load) == np.shape(pv), "负载与光伏形状不一致"
    assert np.shape(load)[-1] == len(price), "电价长度不一致"
    assert cfg.s_min <= s0 <= cfg.s_max, "运行初始SOC越界"


def _rows_to_sparse(rows, cols, vals, rhs, nvars):
    return coo_matrix((vals, (rows, cols)), shape=(len(rhs), nvars)).tocsr()


def solve_plan(load_scen, pv_scen, weights, price, s0, cfg, terminal_free=False, force_milp=False):
    """[M7-M9][S1-S2] 零点计划：先LP松弛，可恢复整数模式即获最优证书，否则HiGHS MILP。"""
    tic = perf_counter()
    L = np.atleast_2d(np.asarray(load_scen, dtype=float))
    V = np.atleast_2d(np.asarray(pv_scen, dtype=float))
    p = np.asarray(price, dtype=float)
    w = np.asarray(weights, dtype=float)
    _check_arrays(L, V, p, s0, cfg)
    K, T = L.shape
    assert len(w) == K and np.all(w >= 0) and np.isclose(w.sum(), 1), "情景概率无效"
    count = K * T
    idx_s = T + np.arange(count).reshape(K, T)
    idx_e = T + count + np.arange(count).reshape(K, T)
    idx_h = T + 2 * count + np.arange(K)
    nv = T + 2 * count + K
    objective = np.zeros(nv)
    objective[:T] = p
    objective[idx_e] = w[:, None] * cfg.emergency_multiplier * p[None, :]
    objective[idx_h] = 0 if terminal_free else w * cfg.kappa
    low = np.zeros(nv)
    high = np.full(nv, np.inf)
    low[idx_s], high[idx_s] = cfg.s_min, cfg.s_max
    high[idx_e] = L
    rr, cc, vv, rhs = [], [], [], []

    def add(coeff, bound):
        r = len(rhs)
        for j, a in coeff:
            rr.append(r); cc.append(int(j)); vv.append(float(a))
        rhs.append(float(bound))

    def delta_coeff(k, t, scale):
        pair = [(idx_s[k, t], scale)]
        if t:
            pair.append((idx_s[k, t - 1], -scale))
        return pair

    cap_c = cfg.eta_c * cfg.p_max * cfg.dt
    cap_d = np.minimum(cfg.p_max * cfg.dt, L) / cfg.eta_d
    for k in range(K):
        for t in range(T):
            for slope in (1 / cfg.eta_c, cfg.eta_d):
                add([(t, -1), (idx_e[k, t], -1)] + delta_coeff(k, t, slope),
                    -(L[k, t] - V[k, t]) + (slope * s0 if t == 0 else 0))
            add(delta_coeff(k, t, 1), cap_c + (s0 if t == 0 else 0))
            add(delta_coeff(k, t, -1), cap_d[k, t] - (s0 if t == 0 else 0))
        add([(idx_s[k, -1], -1), (idx_h[k], -1)], -cfg.s_ref)
    A = _rows_to_sparse(rr, cc, vv, rhs, nv)
    ub = np.asarray(rhs)
    result = linprog(objective, A_ub=A, b_ub=ub, bounds=list(zip(low, high)), method="highs")
    assert result.success, f"日前LP失败：{result.message}"
    state = np.c_[np.full(K, s0), result.x[idx_s]]
    dx = np.diff(state, axis=1)
    emergency = result.x[idx_e]
    relaxed = not ((dx > cfg.energy_tol) & (emergency > cfg.energy_tol)).any()
    status = "LP_CERTIFIED_MILP_OPTIMAL"
    gap = 0.0
    lp_bound = float(result.fun)
    if force_milp or not relaxed:
        idx_z = nv + np.arange(count).reshape(K, T)
        for k in range(K):
            for t in range(T):
                add(delta_coeff(k, t, 1) + [(idx_z[k, t], -cap_c)], s0 if t == 0 else 0)
                add(delta_coeff(k, t, -1) + [(idx_z[k, t], cap_d[k, t])],
                    cap_d[k, t] - (s0 if t == 0 else 0))
                add([(idx_e[k, t], 1), (idx_z[k, t], L[k, t])], L[k, t])
        Am = _rows_to_sparse(rr, cc, vv, rhs, nv + count)
        ub_m = np.asarray(rhs)
        mip_obj = np.r_[objective, np.zeros(count)]
        mip_int = np.r_[np.zeros(nv), np.ones(count)]
        mip_bnd = Bounds(np.r_[low, np.zeros(count)], np.r_[high, np.ones(count)])
        mip_con = LinearConstraint(Am, -np.inf, ub_m)
        result = milp(mip_obj, integrality=mip_int, bounds=mip_bnd, constraints=mip_con,
                      options={"time_limit": cfg.mip_time_limit, "mip_rel_gap": cfg.mip_rel_gap})
        if result.x is None:
            warnings.warn("全情景MILP在时限内未找到可行解，按4倍时限与放宽精度自动重试一次")
            result = milp(mip_obj, integrality=mip_int, bounds=mip_bnd, constraints=mip_con,
                          options={"time_limit": 4.0 * cfg.mip_time_limit,
                                   "mip_rel_gap": max(cfg.mip_rel_gap, 1e-2)})
        if result.x is None:
            warnings.warn("重试仍无可行解：该日改用保守回退计划（逐时段取各情景净负荷最大值），"
                          "因果可执行、成本略偏高，全年运行不中断")
            plan_c = np.maximum(L - V, 0.0).max(axis=0)
            return {"plan": plan_c, "scenario_soc": np.full((K, T + 1), np.nan),
                    "scenario_emergency": np.zeros((K, T)), "scenario_charge": np.zeros((K, T)),
                    "scenario_discharge": np.zeros((K, T)),
                    "status": "FALLBACK_CONSERVATIVE_PLAN", "gap": float("nan"),
                    "objective": float(p @ plan_c), "lp_bound": lp_bound,
                    "relaxation_tight": False, "solve_seconds": perf_counter() - tic}
        assert np.max(Am @ result.x - ub_m) <= cfg.energy_tol, "MILP约束违反"
        gap = float(getattr(result, "mip_gap", np.nan))
        status = "MILP_OPTIMAL" if result.success else "MILP_FEASIBLE_TIME_LIMIT"
        if not result.success:
            warnings.warn(f"MILP仅取得可行解，gap={gap:.3g}，不会宣称全局最优")
        state = np.c_[np.full(K, s0), result.x[idx_s]]
        dx = np.diff(state, axis=1)
        emergency = result.x[idx_e]
    xvec = result.x[:nv]
    assert np.max(A @ xvec - ub) <= cfg.energy_tol, "日前基础约束违反"
    assert not ((dx > cfg.energy_tol) & (emergency > cfg.energy_tol)).any(), "禁止应急购电充电"
    charge = np.maximum(dx, 0) / cfg.eta_c
    discharge = cfg.eta_d * np.maximum(-dx, 0)
    plan = np.maximum(xvec[:T], 0)
    assert np.max(charge) <= cfg.p_max * cfg.dt + cfg.energy_tol
    assert np.max(discharge) <= cfg.p_max * cfg.dt + cfg.energy_tol
    supply_margin = plan[None, :] + V + discharge + emergency - L - charge
    assert supply_margin.min() >= -cfg.energy_tol, "日前电量不足"
    return {"plan": plan, "scenario_soc": state, "scenario_emergency": emergency,
            "scenario_charge": charge, "scenario_discharge": discharge, "status": status,
            "gap": gap, "objective": float(objective @ xvec), "lp_bound": lp_bound,
            "relaxation_tight": relaxed, "solve_seconds": perf_counter() - tic}


# ===================== 4. [RT] 冻结计划下的因果实时回放 =====================
def _mpc_first_action(net_forecast, price, plan, s0, cfg, terminal_free=False):
    """[RT1-RT3] 固定 q 后的剩余日内LP；只返回当前一步动作。"""
    n = np.asarray(net_forecast)
    H = len(n)
    surplus = np.maximum(plan - n, 0)
    deficit = np.maximum(n - plan, 0)
    upper_dx = cfg.eta_c * np.minimum(cfg.p_max * cfg.dt, surplus)
    lower_dx = -np.minimum(cfg.p_max * cfg.dt, deficit) / cfg.eta_d
    D = diags([np.ones(H), -np.ones(max(H - 1, 0))], [0, -1], shape=(H, H), format="csr")
    init = np.zeros(H); init[0] = s0
    marginal = cfg.emergency_multiplier * np.asarray(price) * cfg.eta_d * (deficit > 0)
    cost_s = np.asarray(D.T @ marginal).ravel()
    cost = np.r_[cost_s, 0 if terminal_free else cfg.kappa]
    Z = csr_matrix((H, 1))
    terminal = csr_matrix(([-1., -1.], ([0, 0], [H - 1, H])), shape=(1, H + 1))
    A = vstack([hstack([D, Z]), hstack([-D, Z]), terminal], format="csr")
    b = np.r_[upper_dx + init, -lower_dx - init, -cfg.s_ref]
    result = linprog(cost, A_ub=A, b_ub=b,
                     bounds=[(cfg.s_min, cfg.s_max)] * H + [(0, None)], method="highs")
    assert result.success, f"实时LP失败：{result.message}"
    dx = result.x[0] - s0
    return max(dx, 0) / cfg.eta_c, cfg.eta_d * max(-dx, 0)


def check_dispatch(frame, cfg):
    """[QA1] 独立物理检查已执行轨迹，不依赖优化器返回状态。"""
    tol = cfg.energy_tol
    L, V, q = (frame[k].to_numpy() for k in ["load_kwh", "pv_kwh", "plan_kwh"])
    c, d, e = (frame[k].to_numpy() for k in ["charge_kwh", "discharge_kwh", "emergency_kwh"])
    a, u = (frame[k].to_numpy() for k in ["plan_used_kwh", "pv_used_kwh"])
    s0, s1 = (frame[k].to_numpy() for k in ["soc_start_kwh", "soc_end_kwh"])
    residual = a + u + d + e - L - c
    soc_residual = s1 - s0 - cfg.eta_c * c + d / cfg.eta_d
    assert np.max(np.abs(residual)) <= tol, "执行轨迹电量不守恒"
    assert np.max(np.abs(soc_residual)) <= tol, "SOC递推错误"
    assert min(s0.min(), s1.min()) >= cfg.s_min - tol
    assert max(s0.max(), s1.max()) <= cfg.s_max + tol
    assert np.all(c >= -tol) and np.all(d >= -tol) and np.all(e >= -tol)
    assert max(c.max(), d.max()) <= cfg.p_max * cfg.dt + tol
    assert not ((c > tol) & (d > tol)).any(), "同时充放电"
    assert not ((c > tol) & (e > tol)).any(), "应急购电用于充电"
    assert np.all(a >= -tol) and np.all(a <= q + tol)
    assert np.all(u >= -tol) and np.all(u <= V + tol)
    assert np.max(np.abs(e - np.maximum(L - V - q - d, 0))) <= tol, "应急量不是实际剩余负荷缺口"
    p = frame.price.to_numpy()
    assert np.allclose(frame.total_cost, p * q + cfg.emergency_multiplier * p * e, atol=tol, rtol=1e-10)
    if len(frame) > 1:
        assert np.max(np.abs(s1[:-1] - s0[1:])) <= tol, "储电量不连续"
    return {"balance_max_abs_kwh": float(abs(residual).max()),
            "soc_max_abs_kwh": float(abs(soc_residual).max()), "passed": True}


def simulate_day(actual_load, actual_pv, forecast_load, forecast_pv, price, plan, s0, cfg,
                 terminal_free=False):
    """[RT0-RT4] 当前段用实测、未来段只用零点预测；计划 q 冻结，只纠偏储能与紧急购电。"""
    L, V, p, q = map(lambda a: np.asarray(a, dtype=float).copy(), (actual_load, actual_pv, price, plan))
    Lh, Vh = map(lambda a: np.asarray(a, dtype=float).copy(), (forecast_load, forecast_pv))
    _check_arrays(L, V, p, s0, cfg)
    assert np.isfinite(q).all() and np.all(q >= 0) and len(q) == len(L), "计划量无效"
    assert Lh.shape == L.shape and Vh.shape == V.shape and np.isfinite(Lh + Vh).all(), "预测形状无效"
    q_frozen = q.copy()
    state, records = float(s0), []
    for t in range(len(L)):
        n_future = (Lh[t:] - Vh[t:]).copy()
        n_future[0] = L[t] - V[t]
        c, d = _mpc_first_action(n_future, p[t:], q[t:], state, cfg, terminal_free)
        surplus = max(q[t] + V[t] - L[t], 0)
        deficit = max(L[t] - V[t] - q[t], 0)
        c = min(max(c, 0), cfg.p_max * cfg.dt, surplus, max((cfg.s_max - state) / cfg.eta_c, 0))
        d = min(max(d, 0), cfg.p_max * cfg.dt, deficit, max((state - cfg.s_min) * cfg.eta_d, 0))
        e = max(deficit - d, 0)
        demand = L[t] + c - d - e
        u = min(V[t], max(demand, 0))
        a = max(demand - u, 0)
        next_state = state + cfg.eta_c * c - d / cfg.eta_d
        start, end = t * 10, (t + 1) * 10
        records.append(dict(
            t=t, interval=f"{start // 60:02d}:{start % 60:02d}-{end // 60:02d}:{end % 60:02d}",
            price=p[t], load_kwh=L[t], pv_kwh=V[t], load_hat_kwh=Lh[t], pv_hat_kwh=Vh[t],
            plan_kwh=q[t], charge_kwh=c, discharge_kwh=d, emergency_kwh=e,
            pv_used_kwh=u, plan_used_kwh=a, curtailment_kwh=max(V[t] - u, 0),
            unused_plan_kwh=max(q[t] - a, 0), soc_start_kwh=state, soc_end_kwh=next_state,
            plan_cost=p[t] * q[t], emergency_cost=cfg.emergency_multiplier * p[t] * e,
            total_cost=p[t] * (q[t] + cfg.emergency_multiplier * e)))
        state = next_state
    assert np.array_equal(q, q_frozen), "日内修改了零点计划"
    frame = pd.DataFrame(records)
    frame.attrs["checks"] = check_dispatch(frame, cfg)
    return frame


def summarize_day(frame):
    keys = ["plan_kwh", "emergency_kwh", "plan_cost", "emergency_cost", "total_cost",
            "charge_kwh", "discharge_kwh", "curtailment_kwh", "unused_plan_kwh"]
    row = {k: float(frame[k].sum()) for k in keys}
    row.update(soc_start_kwh=float(frame.soc_start_kwh.iloc[0]),
               soc_end_kwh=float(frame.soc_end_kwh.iloc[-1]),
               emergency_intervals=int((frame.emergency_kwh > 2e-5).sum()))
    return row


def run_period(data, forecast, cfg, start_day, end_day, method="stochastic", s0=None, progress=True):
    """[R0] 跨日滚动：逐日零点求解→冻结计划→因果回放→SOC自然传递。"""
    state = cfg.s_initial if s0 is None else float(s0)
    frames, summaries, audits = [], [], []
    for day in range(start_day, end_day):
        date = pd.Timestamp(data["dates"][day])
        if method == "point":
            Ls, Vs = forecast["load_hat"][day:day + 1], forecast["pv_hat"][day:day + 1]
            weights, sources = np.ones(1), np.array([], dtype=int)
        else:
            Ls, Vs, weights, sources = make_scenarios(
                day, forecast, data["load"], data["pv"], scenario_count=cfg.scenarios,
                history_days=cfg.history_days, seed=cfg.seed, residual_scale=cfg.residual_scale)
        assert len(sources) == 0 or max(sources) < day, "场景包含当日或未来误差"
        free = day == len(data["dates"]) - 1
        solution = solve_plan(Ls, Vs, weights, data["price"], state, cfg, terminal_free=free)
        frame = simulate_day(data["load"][day], data["pv"][day], forecast["load_hat"][day],
                             forecast["pv_hat"][day], data["price"], solution["plan"], state, cfg,
                             terminal_free=free)
        frame.insert(0, "date", date)
        summary = summarize_day(frame)
        summary.update(date=date, method=method, solve_seconds=solution["solve_seconds"],
                       solve_status=solution["status"], mip_gap=solution["gap"],
                       surrogate_objective=solution["objective"],
                       source_day_max=int(max(sources)) if len(sources) else day - 1)
        audits.append(dict(date=date, method=method, **frame.attrs["checks"],
                           plan_frozen=True, source_causal=True,
                           solve_status=solution["status"], gap=solution["gap"]))
        state = summary["soc_end_kwh"]
        frames.append(frame); summaries.append(summary)
        if progress and (day == start_day or (day - start_day + 1) % 30 == 0 or day == end_day - 1):
            print(f"{method} {date.date()} | cost={summary['total_cost']:,.2f} | "
                  f"emergency={summary['emergency_kwh']:,.2f} kWh | {solution['status']}", flush=True)
    detail = pd.concat(frames, ignore_index=True)
    daily = pd.DataFrame(summaries)
    audit = pd.DataFrame(audits)
    assert np.allclose(daily.soc_end_kwh.to_numpy()[:-1], daily.soc_start_kwh.to_numpy()[1:],
                       atol=cfg.energy_tol), "跨日SOC不连续"
    return detail, daily, audit


# ===================== 5. [R] 结果导出（result2.xlsx 等） =====================
def _clock(t: int) -> str:
    return f"{t // 6:02d}:{(t % 6) * 10:02d}"


def _events(detail):
    """[R3] 日内连续 e_t>EPS 时段合并为事件；无事件日期明确填0。"""
    rows = []
    for date, block in detail.groupby("date", sort=True):
        block = block.sort_values("t").reset_index(drop=True)
        active = block["emergency_kwh"].to_numpy() > EPS
        starts = np.flatnonzero(active & ~np.r_[False, active[:-1]])
        ends = np.flatnonzero(active & ~np.r_[active[1:], False])
        for start, end in zip(starts, ends):
            event = block.iloc[start:end + 1]
            first, last = int(event["t"].iloc[0]), int(event["t"].iloc[-1])
            rows.append({"日期": date, "购电时间段": f"{_clock(first)}-{_clock(last + 1)}",
                         "购电量": float(event["emergency_kwh"].sum()),
                         "紧急购电费(元)": float(event["emergency_cost"].sum()),
                         "说明": "连续10分钟时段合并；电价分段累计"})
        if not len(starts):
            rows.append({"日期": date, "购电时间段": "无紧急购电", "购电量": 0.0,
                         "紧急购电费(元)": 0.0, "说明": "该日所有时段紧急购电量均不超过数值容差"})
    return pd.DataFrame(rows, columns=["日期", "购电时间段", "购电量", "紧急购电费(元)", "说明"])


def _blocks(detail):
    """[R1,R2] 每日6个四小时块；SOC只取日初/日末，不相加。"""
    records = []
    for date, day in detail.groupby("date", sort=True):
        day = day.sort_values("t")
        for number in range(6):
            part = day.loc[(day["t"] >= number * 24) & (day["t"] < (number + 1) * 24)]
            records.append({"日期": date, "时间段": f"{number * 4:02d}:00-{(number + 1) * 4:02d}:00",
                            "充电量": float(part["charge_kwh"].sum()),
                            "放电量": float(part["discharge_kwh"].sum()),
                            "时刻": "0:00" if number == 0 else ("24:00" if number == 1 else None),
                            "储电量": float(day["soc_start_kwh"].iloc[0]) if number == 0 else
                            (float(day["soc_end_kwh"].iloc[-1]) if number == 1 else None)})
    return pd.DataFrame(records)


def _write_sheet(wb, name, header, rows):
    ws = wb.create_sheet(name)
    ws.append([str(h) for h in header])
    for r in rows:
        ws.append([v.to_pydatetime() if isinstance(v, (pd.Timestamp, datetime)) else
                   (float(v) if isinstance(v, np.floating) else
                    (int(v) if isinstance(v, np.integer) else v)) for v in r])
    return ws


def _frame_rows(frame):
    return [list(r) for r in frame.itertuples(index=False, name=None)]


def export_results(detail, daily, audit, comparison, time_map, out_dir: Path, metadata: dict):
    """[R1-R5] 导出 result2.xlsx（计划购电量/充放电量/紧急购电量/指定日期/时间映射/口径）与CSV。"""
    out_dir = Path(out_dir); out_dir.mkdir(parents=True, exist_ok=True)
    detail = detail.copy(); daily = daily.copy()
    detail["date"] = pd.to_datetime(detail["date"]).dt.normalize()
    daily["date"] = pd.to_datetime(daily["date"]).dt.normalize()
    detail = detail.sort_values(["date", "t"]).reset_index(drop=True)
    daily = daily.sort_values("date").reset_index(drop=True)
    assert not detail.duplicated(["date", "t"]).any(), "存在重复日期/时段键"
    assert detail.groupby("date").size().eq(144).all(), "每天必须144段"
    assert np.allclose(detail["total_cost"], detail["plan_cost"] + detail["emergency_cost"], atol=1e-5), "实际费用组成错误"
    q = detail.pivot(index="date", columns="t", values="plan_kwh").reindex(columns=range(144))
    assert q.notna().all().all(), "购电矩阵存在缺失时段"
    dates = pd.DatetimeIndex(q.index)
    grouped = detail.groupby("date", sort=True)
    totals = grouped[["plan_kwh", "plan_cost", "emergency_cost", "total_cost", "emergency_kwh"]].sum()
    intervals = [f"{_clock(t)}-{_clock(t + 1)}" for t in range(144)]

    header, plan_rows = None, []
    if USE_TEMPLATE_HEADER and RESULT_TEMPLATE.is_file():   # 逐字复制官方模板147列表头
        wb_t = load_workbook(RESULT_TEMPLATE, data_only=True, read_only=True)
        try:
            header = [c for c in next(wb_t["计划购电量"].iter_rows(min_row=1, max_row=1, values_only=True))]
        finally:
            wb_t.close()
        assert len(header) == 147, f"官方模板表头应为147列，实际{len(header)}列"
        for date in dates:
            v = totals.loc[date]
            plan_rows.append([date] + [float(x) for x in q.loc[date]] +
                             [float(v["plan_kwh"]), float(v["total_cost"])])
    else:                                                    # 规范区间表头 + 费用细分列
        header = ["日期\\时间"] + intervals + ["全天购电量", "全天购电费",
                                              "其中计划购电费", "其中紧急购电费", "另列紧急购电量"]
        for date in dates:
            v = totals.loc[date]
            plan_rows.append([date] + [float(x) for x in q.loc[date]] +
                             [float(v["plan_kwh"]), float(v["total_cost"]),
                              float(v["plan_cost"]), float(v["emergency_cost"]),
                              float(v["emergency_kwh"])])

    wb = Workbook(); wb.remove(wb.active)
    _write_sheet(wb, "计划购电量", header, plan_rows)
    _write_sheet(wb, "充放电量", list(_blocks(detail).columns), _frame_rows(_blocks(detail)))
    _write_sheet(wb, "紧急购电量", list(_events(detail).columns), _frame_rows(_events(detail)))

    special_rows, special_header = [], None
    for date in pd.to_datetime(list(SPECIAL_DATES)):
        if date not in set(dates):
            continue
        day = detail[detail["date"] == date].set_index("t")
        if special_header is None:
            special_header = ["日期"] + [f"{intervals[t]}计划购电量(kWh)" for t in (60, 72, 84, 96, 108, 120)] + \
                             ["全天计划购电量(kWh)", "全天实际购电费(元)", "其中计划购电费(元)", "其中紧急购电费(元)"]
        v = totals.loc[date]
        special_rows.append([date] + [float(day.loc[t, "plan_kwh"]) for t in (60, 72, 84, 96, 108, 120)] +
                            [float(v["plan_kwh"]), float(v["total_cost"]),
                             float(v["plan_cost"]), float(v["emergency_cost"])])
    if special_header:
        _write_sheet(wb, "指定日期结果", special_header, special_rows)
    _write_sheet(wb, "时间映射", list(time_map.columns), _frame_rows(time_map))
    meta = pd.DataFrame([{"参数": str(k), "参数值": str(v)} for k, v in metadata.items()])
    _write_sheet(wb, "口径与参数", ["参数", "参数值"], _frame_rows(meta))
    result_path = out_dir / "result2.xlsx"
    wb.save(result_path)

    detail.to_csv(out_dir / "逐段明细.csv", index=False, encoding="utf-8-sig")
    daily.to_csv(out_dir / "逐日汇总.csv", index=False, encoding="utf-8-sig")
    audit.to_csv(out_dir / "物理审计.csv", index=False, encoding="utf-8-sig")
    comparison.to_csv(out_dir / "方案对比.csv", index=False, encoding="utf-8-sig")
    print(f"[R] 已导出 {result_path}：{len(dates)}天×144段；事件{len(_events(detail))}行。")
    return result_path


# ===================== 6. 可选：DeepSeek API 生成结果分析文字 =====================
def deepseek_analysis(prompt_text: str):
    """仅当 USE_DEEPSEEK_REPORT=True 且填写了 Key 时调用；失败只警告不中断。"""
    if not (USE_DEEPSEEK_REPORT and DEEPSEEK_API_KEY):
        return None
    body = json.dumps({"model": DEEPSEEK_MODEL, "temperature": 0.3,
                       "messages": [{"role": "system", "content": "你是数学建模论文写作助手，用中文作答。"},
                                    {"role": "user", "content": prompt_text}]}).encode("utf-8")
    req = urllib.request.Request(DEEPSEEK_BASE_URL.rstrip("/") + "/chat/completions", data=body,
                                 headers={"Content-Type": "application/json",
                                          "Authorization": f"Bearer {DEEPSEEK_API_KEY}"})
    try:
        with urllib.request.urlopen(req, timeout=180) as resp:
            return json.loads(resp.read().decode("utf-8"))["choices"][0]["message"]["content"]
    except Exception as exc:
        warnings.warn(f"DeepSeek API 调用失败，已跳过分析生成：{exc}")
        return None


# ===================== 7. 主流程 =====================
def main():
    cfg = Config().validate()
    data = load_inputs(INPUT_DIR)
    print(data["input_audit"].to_string(index=False))
    forecast = prepare_forecasts(data["load"], data["pv"], window=cfg.window)
    err = point_summary(forecast, data["load"], data["pv"], data["dates"], start_day=31)
    print("[F2] 2月起平均误差(kWh)：load_MAE=%.2f pv_MAE=%.2f net_MAE=%.2f" % (
        err.load_MAE_kWh.mean(), err.pv_MAE_kWh.mean(), err.net_MAE_kWh.mean()))

    start_day = 0 if JAN_WARMUP else 31
    s0 = None if JAN_WARMUP else cfg.s_initial
    print(f"[RUN] 起始日索引={start_day}（{'一月连续预热' if JAN_WARMUP else '2月1日重置6000'}）")
    detail_s, daily_s, audit_s = run_period(data, forecast, cfg, start_day, 365, "stochastic", s0)
    detail_p, daily_p, audit_p = run_period(data, forecast, cfg, start_day, 365, "point", s0)

    cut = pd.Timestamp(OUTPUT_START)
    detail = detail_s[detail_s["date"] >= cut].reset_index(drop=True)
    daily = daily_s[daily_s["date"] >= cut].reset_index(drop=True)
    audit = audit_s[audit_s["date"] >= cut].reset_index(drop=True)
    daily_p = daily_p[daily_p["date"] >= cut].reset_index(drop=True)

    def _cmp(tag, d):
        return {"方案": tag, "总购电费(元)": float(d.total_cost.sum()),
                "计划购电费(元)": float(d.plan_cost.sum()), "紧急购电费(元)": float(d.emergency_cost.sum()),
                "计划购电量(kWh)": float(d.plan_kwh.sum()), "紧急购电量(kWh)": float(d.emergency_kwh.sum()),
                "紧急天数": int((d.groupby('date').emergency_kwh.sum() > 2e-5).sum()),
                "弃光量(kWh)": float(d.curtailment_kwh.sum()), "期末SOC(kWh)": float(d.soc_end_kwh.iloc[-1])}
    comparison = pd.DataFrame([_cmp("联合误差情景(主方案)", daily), _cmp("点预测对照", daily_p)])
    print(comparison.to_string(index=False))

    metadata = {
        "输出期": f"{OUTPUT_START}—2025-12-31，每天144段",
        "一月处理": "1月1日起连续因果仿真，SOC自然传递" if JAN_WARMUP else "2月1日重置为6000 kWh（声明的初始化假设）",
        "时间映射": "源标签0:10为区间终点，对应物理区间00:00-00:10；0:00+1对应23:50-24:00。",
        "模板说明": "官方模板表头自0:10-0:20起、整体后错一段；本文件按USE_TEMPLATE_HEADER设置输出，时间映射表保留对应关系。",
        "费用口径": "实际费用=Σp_t q_t+Σ5p_t e_t；终端SOC惩罚不计入。",
        "计划电弃用": "计划电允许未取用但全额付费；未用计划量与弃光分别统计。",
        "紧急事件": "连续e_t>1e-7时段合并展示，费用逐段计价；无事件日期填0。",
        "参数": json.dumps(asdict(cfg), ensure_ascii=False),
        "输入文件": json.dumps(data["source_paths"], ensure_ascii=False),
    }
    out_path = export_results(detail, daily, audit, comparison, data["time_map"], OUTPUT_DIR, metadata)

    text = deepseek_analysis(
        "以下是微电网第二问全年回放的汇总结果，请写一段300字以内的论文结果分析：\n" +
        comparison.to_string(index=False) + "\n指定日期：\n" +
        detail[detail["date"].isin(pd.to_datetime(list(SPECIAL_DATES)))].groupby("date")[
            ["plan_kwh", "emergency_kwh", "total_cost"]].sum().to_string())
    if text:
        (OUTPUT_DIR / "结果分析_deepseek.md").write_text(text, encoding="utf-8")
        print("[DeepSeek] 结果分析已保存。")
    print(f"[DONE] 主结果：{out_path}")


if __name__ == "__main__":
    main()
