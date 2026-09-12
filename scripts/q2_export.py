"""Question 2 CSV outputs and strict official-template workbook export."""
from __future__ import annotations

from copy import deepcopy
from io import BytesIO
import os
from pathlib import Path, PurePosixPath
import tempfile
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

from project_paths import TEMPLATE_RESULT_2


SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"main": SPREADSHEET_NS, "rel": DOCUMENT_REL_NS, "pkg": PACKAGE_REL_NS}
READBACK_TOLERANCE = 1.0e-6
DISPLAY_ZERO_TOLERANCE = 1.0e-8
BLOCK_LABELS = (
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
)


def _display_float(value: float) -> float:
    value = float(value)
    return 0.0 if abs(value) < DISPLAY_ZERO_TOLERANCE else value


def four_hour_table(detail: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for date, day in detail.groupby("date", sort=True):
        day = day.sort_values("t")
        for block, label in enumerate(BLOCK_LABELS):
            part = day[(day["t"] >= block * 24) & (day["t"] < (block + 1) * 24)]
            rows.append(
                {
                    "date": pd.Timestamp(date).normalize(),
                    "block": block,
                    "time_block": label,
                    "charge_kwh": float(part["charge_kwh"].sum()),
                    "discharge_kwh": float(part["discharge_kwh"].sum()),
                    "soc_time": "0:00" if block == 0 else ("24:00" if block == 1 else ""),
                    "soc_kwh": (
                        float(day["soc_start_kwh"].iloc[0])
                        if block == 0
                        else (float(day["soc_end_kwh"].iloc[-1]) if block == 1 else np.nan)
                    ),
                }
            )
    return pd.DataFrame(rows)


def emergency_event_table(detail: pd.DataFrame, tolerance: float) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for date, day in detail.groupby("date", sort=True):
        day = day.sort_values("t").reset_index(drop=True)
        active = day["emergency_kwh"].to_numpy() > tolerance
        starts = np.flatnonzero(active & ~np.r_[False, active[:-1]])
        ends = np.flatnonzero(active & ~np.r_[active[1:], False])
        if len(starts) == 0:
            rows.append(
                {
                    "date": pd.Timestamp(date).normalize(),
                    "event_index": 0,
                    "purchase_interval": "无紧急购电",
                    "emergency_kwh": 0.0,
                }
            )
            continue
        for event_index, (start, end) in enumerate(zip(starts, ends)):
            first = int(day.loc[start, "t"])
            last = int(day.loc[end, "t"])
            start_label = f"{first // 6:02d}:{(first % 6) * 10:02d}"
            end_minutes = (last + 1) * 10
            end_label = (
                "24:00"
                if end_minutes == 24 * 60
                else f"{end_minutes // 60:02d}:{end_minutes % 60:02d}"
            )
            rows.append(
                {
                    "date": pd.Timestamp(date).normalize(),
                    "event_index": event_index,
                    "purchase_interval": f"{start_label}-{end_label}",
                    "emergency_kwh": float(
                        day.loc[start:end, "emergency_kwh"].sum()
                    ),
                }
            )
    return pd.DataFrame(rows)


def _register_xml_namespaces(xml_bytes: bytes) -> None:
    for _, namespace in ET.iterparse(BytesIO(xml_bytes), events=("start-ns",)):
        prefix, uri = namespace
        ET.register_namespace(prefix or "", uri)


def _worksheet_parts(archive: ZipFile) -> dict[str, str]:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        node.attrib["Id"]: node.attrib["Target"]
        for node in relationships_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    parts: dict[str, str] = {}
    for sheet in workbook_root.findall("main:sheets/main:sheet", NS):
        relationship_id = sheet.attrib[f"{{{DOCUMENT_REL_NS}}}id"]
        target = relationships[relationship_id]
        if target.startswith("/"):
            part = target.lstrip("/")
        else:
            part = str(PurePosixPath("xl") / PurePosixPath(target))
        parts[sheet.attrib["name"]] = part
    return parts


def _column_number(cell_reference: str) -> int:
    value = 0
    for character in cell_reference:
        if not character.isalpha():
            break
        value = value * 26 + ord(character.upper()) - ord("A") + 1
    return value


def _find_or_create_cell(row: ET.Element, reference: str, style: str | None = None) -> ET.Element:
    cell = row.find(f"main:c[@r='{reference}']", NS)
    if cell is None:
        attributes = {"r": reference}
        if style is not None:
            attributes["s"] = style
        cell = ET.Element(f"{{{SPREADSHEET_NS}}}c", attributes)
        target_column = _column_number(reference)
        insert_at = len(row)
        for index, existing in enumerate(row.findall("main:c", NS)):
            if _column_number(existing.attrib["r"]) > target_column:
                insert_at = index
                break
        row.insert(insert_at, cell)
    elif style is not None:
        cell.attrib["s"] = style
    return cell


def _clear_cell(cell: ET.Element) -> None:
    cell.attrib.pop("t", None)
    for child in list(cell):
        cell.remove(child)


def _set_numeric_cell(row: ET.Element, reference: str, value: float, style: str | None = None) -> None:
    cell = _find_or_create_cell(row, reference, style)
    _clear_cell(cell)
    value_node = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}v")
    value_node.text = repr(_display_float(value))


def _set_inline_string_cell(row: ET.Element, reference: str, value: str, style: str | None = None) -> None:
    cell = _find_or_create_cell(row, reference, style)
    _clear_cell(cell)
    cell.attrib["t"] = "inlineStr"
    inline = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}is")
    text = ET.SubElement(inline, f"{{{SPREADSHEET_NS}}}t")
    text.text = value


def _set_blank_cell(row: ET.Element, reference: str) -> None:
    cell = _find_or_create_cell(row, reference)
    _clear_cell(cell)


def _rebase_row(prototype: ET.Element, row_number: int) -> ET.Element:
    row = deepcopy(prototype)
    row.attrib["r"] = str(row_number)
    for cell in row.findall("main:c", NS):
        column = "".join(character for character in cell.attrib["r"] if character.isalpha())
        cell.attrib["r"] = f"{column}{row_number}"
    return row


def _set_dimension(root: ET.Element, reference: str) -> None:
    dimension = root.find("main:dimension", NS)
    if dimension is None:
        raise AssertionError("Template worksheet has no dimension element")
    dimension.attrib["ref"] = reference


def _patch_plan_sheet(xml_bytes: bytes, detail: pd.DataFrame) -> bytes:
    _register_xml_namespaces(xml_bytes)
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise AssertionError("Plan worksheet has no sheetData")
    plan = detail.pivot(index="date", columns="t", values="plan_kwh").reindex(columns=range(144))
    totals = detail.groupby("date", sort=True)[["plan_kwh", "total_cost_yuan"]].sum()
    if plan.shape != (334, 144) or plan.isna().any().any():
        raise AssertionError("Plan export requires exactly 334 complete days")
    rows = sheet_data.findall("main:row", NS)
    if len(rows) != 335:
        raise AssertionError("Official plan template does not contain 334 date rows")
    for offset, date in enumerate(plan.index, start=2):
        row = sheet_data.find(f"main:row[@r='{offset}']", NS)
        if row is None:
            raise AssertionError(f"Plan template is missing row {offset}")
        for interval, value in enumerate(plan.loc[date].to_numpy(), start=2):
            _set_numeric_cell(row, f"{get_column_letter(interval)}{offset}", value, "2")
        _set_numeric_cell(row, f"EP{offset}", totals.loc[date, "plan_kwh"], "2")
        _set_numeric_cell(row, f"EQ{offset}", totals.loc[date, "total_cost_yuan"], "2")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _patch_storage_sheet(xml_bytes: bytes, blocks: pd.DataFrame) -> bytes:
    _register_xml_namespaces(xml_bytes)
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise AssertionError("Storage worksheet has no sheetData")
    original_rows = sheet_data.findall("main:row", NS)
    if len(original_rows) < 7:
        raise AssertionError("Storage template lacks six prototype rows")
    header = deepcopy(original_rows[0])
    prototypes = [deepcopy(row) for row in original_rows[1:7]]
    for row in list(sheet_data):
        sheet_data.remove(row)
    sheet_data.append(header)

    row_number = 2
    for record in blocks.itertuples(index=False):
        block = int(record.block)
        row = _rebase_row(prototypes[block], row_number)
        if block == 0:
            _set_numeric_cell(row, f"A{row_number}", to_excel(record.date.to_pydatetime()))
        else:
            _set_blank_cell(row, f"A{row_number}")
        _set_numeric_cell(row, f"C{row_number}", record.charge_kwh)
        _set_numeric_cell(row, f"D{row_number}", record.discharge_kwh)
        if block == 0:
            _set_numeric_cell(row, f"F{row_number}", record.soc_kwh)
        elif block == 1:
            _set_numeric_cell(row, f"F{row_number}", record.soc_kwh)
        else:
            _set_blank_cell(row, f"F{row_number}")
        sheet_data.append(row)
        row_number += 1
    _set_dimension(root, f"A1:F{row_number - 1}")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _patch_emergency_sheet(xml_bytes: bytes, events: pd.DataFrame) -> bytes:
    _register_xml_namespaces(xml_bytes)
    root = ET.fromstring(xml_bytes)
    sheet_data = root.find("main:sheetData", NS)
    if sheet_data is None:
        raise AssertionError("Emergency worksheet has no sheetData")
    original_rows = sheet_data.findall("main:row", NS)
    if len(original_rows) < 4:
        raise AssertionError("Emergency template lacks row prototypes")
    header = deepcopy(original_rows[0])
    prototypes = [deepcopy(row) for row in original_rows[1:4]]
    for row in list(sheet_data):
        sheet_data.remove(row)
    sheet_data.append(header)

    row_number = 2
    for _, group in events.groupby("date", sort=True):
        records = list(group.itertuples(index=False))
        for index, record in enumerate(records):
            if len(records) == 1:
                prototype = prototypes[0]
            elif index == 0:
                prototype = prototypes[0]
            elif index == len(records) - 1:
                prototype = prototypes[2]
            else:
                prototype = prototypes[1]
            row = _rebase_row(prototype, row_number)
            if index == 0:
                _set_numeric_cell(row, f"A{row_number}", to_excel(record.date.to_pydatetime()))
            else:
                _set_blank_cell(row, f"A{row_number}")
            _set_inline_string_cell(row, f"B{row_number}", record.purchase_interval)
            _set_numeric_cell(row, f"C{row_number}", record.emergency_kwh)
            sheet_data.append(row)
            row_number += 1
    _set_dimension(root, f"A1:C{row_number - 1}")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def write_template_result2(
    output_path: Path,
    detail: pd.DataFrame,
    energy_tolerance_kwh: float,
    template_path: Path = TEMPLATE_RESULT_2,
) -> tuple[Path, pd.DataFrame, pd.DataFrame]:
    output_path = Path(output_path)
    template_path = Path(template_path)
    if not template_path.is_file():
        raise FileNotFoundError(f"Official result2 template not found: {template_path}")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    normalized = detail.copy()
    normalized["date"] = pd.to_datetime(normalized["date"]).dt.normalize()
    normalized = normalized.sort_values(["date", "t"]).reset_index(drop=True)
    blocks = four_hour_table(normalized)
    events = emergency_event_table(normalized, energy_tolerance_kwh)

    temp_file = tempfile.NamedTemporaryFile(
        prefix=f"{output_path.stem}-",
        suffix=".xlsx",
        dir=output_path.parent,
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()
    try:
        with ZipFile(template_path, "r") as source, ZipFile(temp_path, "w") as target:
            parts = _worksheet_parts(source)
            replacements = {
                parts["计划购电量"]: lambda data: _patch_plan_sheet(data, normalized),
                parts["充放电量"]: lambda data: _patch_storage_sheet(data, blocks),
                parts["紧急购电量"]: lambda data: _patch_emergency_sheet(data, events),
            }
            target.comment = source.comment
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename in replacements:
                    payload = replacements[item.filename](payload)
                target.writestr(item, payload)
        os.replace(temp_path, output_path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return output_path, blocks, events


def _assert_close(label: str, actual: object, expected: float) -> None:
    if actual is None or abs(float(actual) - float(expected)) > READBACK_TOLERANCE:
        raise AssertionError(f"{label} mismatch: {actual} != {expected}")


def verify_template_result2(
    output_path: Path,
    detail: pd.DataFrame,
    blocks: pd.DataFrame,
    events: pd.DataFrame,
    template_path: Path = TEMPLATE_RESULT_2,
) -> dict[str, object]:
    output_path = Path(output_path)
    template_path = Path(template_path)
    with ZipFile(template_path, "r") as template, ZipFile(output_path, "r") as output:
        template_parts = _worksheet_parts(template)
        output_parts = _worksheet_parts(output)
        if template_parts != output_parts:
            raise AssertionError("result2.xlsx worksheet mapping differs from the template")
        if template.namelist() != output.namelist():
            raise AssertionError("result2.xlsx package members differ from the template")
        worksheet_parts = set(template_parts.values())
        for item in template.infolist():
            if item.filename not in worksheet_parts:
                if template.read(item.filename) != output.read(item.filename):
                    raise AssertionError(f"Unapproved template package part changed: {item.filename}")

    template_workbook = load_workbook(template_path, data_only=False, read_only=False)
    workbook = load_workbook(output_path, data_only=True, read_only=False)
    try:
        if workbook.sheetnames != template_workbook.sheetnames:
            raise AssertionError("result2.xlsx must retain exactly the three template sheets")
        for sheet_name in template_workbook.sheetnames:
            template_header = [cell.value for cell in template_workbook[sheet_name][1]]
            output_header = [cell.value for cell in workbook[sheet_name][1]]
            if output_header != template_header:
                raise AssertionError(f"Header changed in {sheet_name}")

        normalized = detail.copy()
        normalized["date"] = pd.to_datetime(normalized["date"]).dt.normalize()
        plan = normalized.pivot(index="date", columns="t", values="plan_kwh").reindex(columns=range(144))
        totals = normalized.groupby("date", sort=True)[["plan_kwh", "total_cost_yuan"]].sum()
        plan_sheet = workbook["计划购电量"]
        if plan_sheet.max_row != 335 or plan_sheet.max_column != 147:
            raise AssertionError("Plan sheet shape differs from the official format")
        for row_number, date in enumerate(plan.index, start=2):
            if pd.Timestamp(plan_sheet.cell(row_number, 1).value).normalize() != date:
                raise AssertionError(f"Plan sheet date mismatch at row {row_number}")
            for interval, expected in enumerate(plan.loc[date], start=2):
                _assert_close(
                    f"plan row {row_number} interval {interval - 2}",
                    plan_sheet.cell(row_number, interval).value,
                    expected,
                )
            _assert_close(
                f"plan total row {row_number}",
                plan_sheet.cell(row_number, 146).value,
                totals.loc[date, "plan_kwh"],
            )
            _assert_close(
                f"cost total row {row_number}",
                plan_sheet.cell(row_number, 147).value,
                totals.loc[date, "total_cost_yuan"],
            )

        storage_sheet = workbook["充放电量"]
        if storage_sheet.max_row != len(blocks) + 1 or storage_sheet.max_column != 6:
            raise AssertionError("Storage sheet shape is incorrect")
        for row_number, record in enumerate(blocks.itertuples(index=False), start=2):
            expected_date = record.date if record.block == 0 else None
            actual_date = storage_sheet.cell(row_number, 1).value
            if expected_date is None:
                if actual_date is not None:
                    raise AssertionError(f"Unexpected repeated storage date at row {row_number}")
            elif pd.Timestamp(actual_date).normalize() != expected_date:
                raise AssertionError(f"Storage date mismatch at row {row_number}")
            if storage_sheet.cell(row_number, 2).value != record.time_block:
                raise AssertionError(f"Storage block label mismatch at row {row_number}")
            _assert_close("storage charge", storage_sheet.cell(row_number, 3).value, record.charge_kwh)
            _assert_close("storage discharge", storage_sheet.cell(row_number, 4).value, record.discharge_kwh)
            if record.block in (0, 1):
                _assert_close("storage SOC", storage_sheet.cell(row_number, 6).value, record.soc_kwh)

        emergency_sheet = workbook["紧急购电量"]
        if emergency_sheet.max_row != len(events) + 1 or emergency_sheet.max_column != 3:
            raise AssertionError("Emergency sheet shape is incorrect")
        for row_number, record in enumerate(events.itertuples(index=False), start=2):
            expected_date = record.date if record.event_index == 0 else None
            actual_date = emergency_sheet.cell(row_number, 1).value
            if expected_date is None:
                if actual_date is not None:
                    raise AssertionError(f"Unexpected repeated emergency date at row {row_number}")
            elif pd.Timestamp(actual_date).normalize() != expected_date:
                raise AssertionError(f"Emergency date mismatch at row {row_number}")
            if emergency_sheet.cell(row_number, 2).value != record.purchase_interval:
                raise AssertionError(f"Emergency interval mismatch at row {row_number}")
            _assert_close(
                "emergency amount",
                emergency_sheet.cell(row_number, 3).value,
                record.emergency_kwh,
            )
    finally:
        workbook.close()
        template_workbook.close()
    return {
        "passed": True,
        "sheet_names": ["计划购电量", "充放电量", "紧急购电量"],
        "plan_days": 334,
        "storage_rows": len(blocks),
        "emergency_rows": len(events),
        "non_worksheet_package_parts_preserved": True,
    }


def write_mapping_note(path: Path) -> Path:
    text = """# result2.xlsx Mapping Note

The generated `result2.xlsx` retains the official workbook's three worksheets,
column names, formatting resources, and non-worksheet package parts. The
abbreviated storage and emergency examples in the empty template are expanded to
the complete February 1 through December 31 records.

Attachment 2 labels each ten-minute interval by its endpoint. Audit CSV files use
physical intervals `00:00-00:10` through `23:50-24:00`. The official plan sheet
starts at `0:10-0:20` and ends at `0:00-0:10+1`; those supplied labels are not
changed. The 144 chronological planned-purchase values are written in source slot
order, and this known one-slot label offset must be stated when the table is used.

`全天计划购电量` is the sum of the 144 frozen day-ahead purchases. `全天计划购电费`
is the audited actual purchase cost: planned purchase cost plus five-times-price
emergency purchase cost. Emergency rows merge consecutive positive ten-minute
quantities into one event; their amount is summed while cost remains calculated
at the original interval prices in the audit CSV.
"""
    path.write_text(text, encoding="utf-8")
    return path


def export_q2_outputs(
    output_directory: Path,
    detail: pd.DataFrame,
    daily: pd.DataFrame,
    point_daily: pd.DataFrame,
    audit: pd.DataFrame,
    comparison: pd.DataFrame,
    energy_tolerance_kwh: float,
    team_comparison: pd.DataFrame | None = None,
    template_path: Path = TEMPLATE_RESULT_2,
) -> dict[str, Path]:
    output_directory = Path(output_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    paths = {
        "detail": output_directory / "q2_detail.csv",
        "daily": output_directory / "q2_daily.csv",
        "point_daily": output_directory / "q2_point_daily.csv",
        "audit": output_directory / "q2_audit.csv",
        "comparison": output_directory / "q2_comparison.csv",
        "template": output_directory / "result2.xlsx",
        "mapping_note": output_directory / "result2_mapping_note.md",
    }
    detail.to_csv(paths["detail"], index=False, encoding="utf-8-sig", float_format="%.12g")
    daily.to_csv(paths["daily"], index=False, encoding="utf-8-sig", float_format="%.12g")
    point_daily.to_csv(
        paths["point_daily"], index=False, encoding="utf-8-sig", float_format="%.12g"
    )
    audit.to_csv(paths["audit"], index=False, encoding="utf-8-sig", float_format="%.12g")
    comparison.to_csv(
        paths["comparison"], index=False, encoding="utf-8-sig", float_format="%.12g"
    )
    if team_comparison is not None:
        paths["team_comparison"] = output_directory / "q2_team_comparison.csv"
        team_comparison.to_csv(
            paths["team_comparison"],
            index=False,
            encoding="utf-8-sig",
            float_format="%.12g",
        )
    _, blocks, events = write_template_result2(
        paths["template"], detail, energy_tolerance_kwh, template_path
    )
    verify_template_result2(paths["template"], detail, blocks, events, template_path)
    write_mapping_note(paths["mapping_note"])
    return paths
