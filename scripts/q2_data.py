"""Question 2 input loading, causal forecasts, and residual scenarios."""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from project_paths import ATTACHMENT_1, ATTACHMENT_2


INTERVAL_COUNT = 144
DELTA_HOUR = 1.0 / 6.0
YEAR_DAYS = 365


@dataclass(frozen=True)
class Q2SourceData:
    price_yuan_per_kwh: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    dates: tuple[datetime, ...]
    source_endpoints: tuple[str, ...]
    physical_intervals: tuple[str, ...]
    price_path: Path
    actual_path: Path


@dataclass(frozen=True)
class CausalForecast:
    load_hat_kwh: np.ndarray
    pv_hat_kwh: np.ndarray
    load_residual_kwh: np.ndarray
    pv_residual_kwh: np.ndarray
    window_days: int


@dataclass(frozen=True)
class ScenarioSet:
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    weights: np.ndarray
    source_day_indices: np.ndarray


def _clock(total_minutes: int) -> str:
    total_minutes %= 24 * 60
    return f"{total_minutes // 60:02d}:{total_minutes % 60:02d}"


def physical_interval_labels() -> tuple[str, ...]:
    labels = []
    for interval in range(INTERVAL_COUNT):
        start = interval * 10
        end = (interval + 1) * 10
        end_label = "24:00" if end == 24 * 60 else _clock(end)
        labels.append(f"{_clock(start)}-{end_label}")
    return tuple(labels)


def _format_endpoint(value: object, index: int) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if hasattr(value, "hour") and hasattr(value, "minute"):
        if index == INTERVAL_COUNT - 1 and value.hour == 0 and value.minute == 0:
            return "0:00+1"
        return f"{value.hour:02d}:{value.minute:02d}"
    text = str(value).strip()
    return text or ("0:00+1" if index == INTERVAL_COUNT - 1 else _clock((index + 1) * 10))


def _read_actual_sheet(sheet) -> tuple[np.ndarray, tuple[datetime, ...]]:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) != YEAR_DAYS + 1 or len(rows[0]) != INTERVAL_COUNT + 1:
        raise ValueError(
            f"{sheet.title} must contain one header row and 365 x 144 values"
        )
    dates = tuple(row[0] for row in rows[1:])
    if not all(isinstance(value, datetime) for value in dates):
        raise ValueError(f"{sheet.title} contains a non-datetime date")
    values_kw = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    return values_kw * DELTA_HOUR, dates


def load_q2_source(
    price_path: Path = ATTACHMENT_1,
    actual_path: Path = ATTACHMENT_2,
) -> Q2SourceData:
    price_path = Path(price_path).resolve()
    actual_path = Path(actual_path).resolve()
    if not price_path.is_file():
        raise FileNotFoundError(f"Attachment 1 not found: {price_path}")
    if not actual_path.is_file():
        raise FileNotFoundError(f"Attachment 2 not found: {actual_path}")

    price_workbook = load_workbook(price_path, data_only=True, read_only=True)
    try:
        price_rows = list(price_workbook.worksheets[0].iter_rows(values_only=True))
    finally:
        price_workbook.close()
    if len(price_rows) != INTERVAL_COUNT + 1:
        raise ValueError("Attachment 1 must contain one header and 144 data rows")
    price = np.asarray([row[1] for row in price_rows[1:]], dtype=float)
    endpoints = tuple(
        _format_endpoint(row[0], index) for index, row in enumerate(price_rows[1:])
    )

    actual_workbook = load_workbook(actual_path, data_only=True, read_only=True)
    try:
        load_sheet = next(
            (
                sheet
                for sheet in actual_workbook.worksheets
                if "负荷" in sheet.title or "负载" in sheet.title
            ),
            None,
        )
        pv_sheet = next(
            (sheet for sheet in actual_workbook.worksheets if "光伏" in sheet.title),
            None,
        )
        if load_sheet is None or pv_sheet is None:
            raise ValueError("Attachment 2 is missing load or photovoltaic sheets")
        load_kwh, dates = _read_actual_sheet(load_sheet)
        pv_kwh, pv_dates = _read_actual_sheet(pv_sheet)
    finally:
        actual_workbook.close()

    expected_dates = tuple(
        datetime(2025, 1, 1) + timedelta(days=offset)
        for offset in range(YEAR_DAYS)
    )
    if dates != pv_dates or dates != expected_dates:
        raise ValueError("Attachment 2 dates must cover every day of 2025 in order")
    for name, values in (
        ("price", price),
        ("load", load_kwh),
        ("photovoltaic", pv_kwh),
    ):
        if not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"{name} contains a non-finite or negative value")
    if price.shape != (INTERVAL_COUNT,) or np.any(price <= 0):
        raise ValueError("Price must contain 144 strictly positive values")

    return Q2SourceData(
        price_yuan_per_kwh=price,
        load_kwh=load_kwh,
        pv_kwh=pv_kwh,
        dates=dates,
        source_endpoints=endpoints,
        physical_intervals=physical_interval_labels(),
        price_path=price_path,
        actual_path=actual_path,
    )


def prepare_causal_forecast(
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    window_days: int = 7,
) -> CausalForecast:
    load = np.asarray(load_kwh, dtype=float)
    pv = np.asarray(pv_kwh, dtype=float)
    if load.shape != pv.shape or load.ndim != 2:
        raise ValueError("Load and photovoltaic arrays must have equal 2-D shapes")
    if window_days < 1:
        raise ValueError("Forecast window must be positive")

    load_hat = np.zeros_like(load)
    pv_hat = np.zeros_like(pv)
    for day in range(1, len(load)):
        start = max(0, day - window_days)
        load_hat[day] = load[start:day].mean(axis=0)
        pv_hat[day] = pv[start:day].mean(axis=0)
    return CausalForecast(
        load_hat_kwh=load_hat,
        pv_hat_kwh=pv_hat,
        load_residual_kwh=load - load_hat,
        pv_residual_kwh=pv - pv_hat,
        window_days=window_days,
    )


def make_joint_residual_scenarios(
    day: int,
    forecast: CausalForecast,
    source: Q2SourceData,
    scenario_count: int = 20,
    history_days: int = 28,
    random_seed: int = 20260911,
    residual_scale: float = 1.0,
) -> ScenarioSet:
    if not 0 <= day < len(source.dates):
        raise IndexError("Target day is outside the source period")
    if scenario_count < 1 or history_days < 1:
        raise ValueError("Scenario count and history length must be positive")
    start = max(forecast.window_days, day - history_days)
    pool = np.arange(start, day, dtype=int)
    if pool.size == 0:
        source_days = np.asarray([], dtype=int)
        scenario_load = forecast.load_hat_kwh[day][None, :].copy()
        scenario_pv = forecast.pv_hat_kwh[day][None, :].copy()
    else:
        generator = np.random.default_rng(random_seed + day)
        count = min(scenario_count, pool.size)
        source_days = np.sort(generator.choice(pool, size=count, replace=False))
        scenario_load = np.maximum(
            0.0,
            forecast.load_hat_kwh[day][None, :]
            + residual_scale * forecast.load_residual_kwh[source_days],
        )
        scenario_pv = np.maximum(
            0.0,
            forecast.pv_hat_kwh[day][None, :]
            + residual_scale * forecast.pv_residual_kwh[source_days],
        )
    weights = np.full(len(scenario_load), 1.0 / len(scenario_load))
    if source_days.size and np.any(source_days >= day):
        raise AssertionError("A residual scenario uses current or future data")
    return ScenarioSet(scenario_load, scenario_pv, weights, source_days)


def point_forecast_scenario(day: int, forecast: CausalForecast) -> ScenarioSet:
    if not 0 <= day < len(forecast.load_hat_kwh):
        raise IndexError("Target day is outside the forecast period")
    return ScenarioSet(
        forecast.load_hat_kwh[day][None, :].copy(),
        forecast.pv_hat_kwh[day][None, :].copy(),
        np.ones(1),
        np.asarray([], dtype=int),
    )
