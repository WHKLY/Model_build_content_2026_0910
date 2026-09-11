"""Question 1 table, workbook, and official-template exports."""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from project_paths import TEMPLATE_RESULT_1
from q1_data import INTERVAL_COUNT, Q1SourceData
from q1_model import BaselineSolution, IntervalData, SOC_INITIAL_KWH, SOC_TERMINAL_KWH, ScheduleSolution
from q1_verify import SolutionCheck

READBACK_TOL = 1.0e-6
DISPLAY_ZERO_TOL = 1.0e-8


def _display_float(value: float) -> float:
    value = float(value)
    return 0.0 if abs(value) < DISPLAY_ZERO_TOL else value


def schedule_table(
    source: Q1SourceData,
    interval_data: IntervalData,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "index": np.arange(1, INTERVAL_COUNT + 1),
            "source_time": source.source_time,
            "interval": source.interval,
            "price_yuan_per_kwh": interval_data.price,
            "load_kw": interval_data.load_kw,
            "pv_forecast_kw": interval_data.pv_kw,
            "load_kwh": interval_data.load_kwh,
            "pv_forecast_kwh": interval_data.pv_kwh,
            "grid_A_kwh": plan_a.grid_kwh,
            "charge_A_kwh": plan_a.charge_kwh,
            "discharge_A_kwh": plan_a.discharge_kwh,
            "curtail_A_kwh": plan_a.curtail_kwh,
            "soc_A_kwh": plan_a.soc_kwh,
            "grid_B_kwh": plan_b.grid_kwh,
            "charge_B_kwh": plan_b.charge_kwh,
            "discharge_B_kwh": plan_b.discharge_kwh,
            "curtail_B_kwh": plan_b.curtail_kwh,
            "soc_B_kwh": plan_b.soc_kwh,
            "diff_grid_B_minus_A_kwh": plan_b.grid_kwh - plan_a.grid_kwh,
            "diff_soc_B_minus_A_kwh": plan_b.soc_kwh - plan_a.soc_kwh,
        }
    )


def summary_table(
    baseline: BaselineSolution,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
    checks: dict[str, SolutionCheck],
) -> pd.DataFrame:
    rows = [
        ("no_storage_cost_yuan", baseline.cost),
        ("no_storage_grid_kwh", baseline.total_grid_kwh),
        ("plan_A_cost_yuan", plan_a.cost),
        ("plan_A_grid_kwh", plan_a.total_grid_kwh),
        ("plan_A_charge_kwh", plan_a.total_charge_kwh),
        ("plan_A_discharge_kwh", plan_a.total_discharge_kwh),
        ("plan_A_curtail_kwh", plan_a.total_curtail_kwh),
        ("plan_A_target_index", plan_a.target_index),
        ("plan_A_target_grid_kwh", plan_a.target_grid_kwh),
        ("plan_A_max_balance_residual", checks["plan_a"].max_balance_residual),
        ("plan_A_max_soc_recurrence_residual", checks["plan_a"].max_soc_recurrence_residual),
        ("plan_A_max_lp_ub_violation", checks["plan_a"].max_lp_ub_violation),
        ("plan_A_max_lp_eq_residual", checks["plan_a"].max_lp_eq_residual),
        ("plan_B_cost_yuan", plan_b.cost),
        ("plan_B_grid_kwh", plan_b.total_grid_kwh),
        ("plan_B_charge_kwh", plan_b.total_charge_kwh),
        ("plan_B_discharge_kwh", plan_b.total_discharge_kwh),
        ("plan_B_curtail_kwh", plan_b.total_curtail_kwh),
        ("plan_B_target_index", plan_b.target_index),
        ("plan_B_target_grid_kwh", plan_b.target_grid_kwh),
        ("plan_B_max_balance_residual", checks["plan_b"].max_balance_residual),
        ("plan_B_max_soc_recurrence_residual", checks["plan_b"].max_soc_recurrence_residual),
        ("plan_B_max_lp_ub_violation", checks["plan_b"].max_lp_ub_violation),
        ("plan_B_max_lp_eq_residual", checks["plan_b"].max_lp_eq_residual),
        ("cost_saving_A_vs_no_storage_yuan", baseline.cost - plan_a.cost),
        ("cost_saving_B_vs_no_storage_yuan", baseline.cost - plan_b.cost),
        ("max_abs_grid_difference_B_minus_A_kwh", float(np.max(np.abs(plan_b.grid_kwh - plan_a.grid_kwh)))),
        ("model_note", "same optimal purchase cost, different interval-level purchase schedules"),
    ]
    return pd.DataFrame(rows, columns=["metric", "value"])


def four_hour_table(plan_a: ScheduleSolution, plan_b: ScheduleSolution | None = None) -> pd.DataFrame:
    labels = ["00:00-04:00", "04:00-08:00", "08:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    rows = []
    for k, label in enumerate(labels):
        idx = slice(k * 24, (k + 1) * 24)
        row = {
            "interval": label,
            "charge_A_kwh": float(np.sum(plan_a.charge_kwh[idx])),
            "discharge_A_kwh": float(np.sum(plan_a.discharge_kwh[idx])),
            "grid_A_kwh": float(np.sum(plan_a.grid_kwh[idx])),
        }
        if plan_b is not None:
            row.update(
                {
                    "charge_B_kwh": float(np.sum(plan_b.charge_kwh[idx])),
                    "discharge_B_kwh": float(np.sum(plan_b.discharge_kwh[idx])),
                    "grid_B_kwh": float(np.sum(plan_b.grid_kwh[idx])),
                }
            )
        rows.append(row)
    return pd.DataFrame(rows)


def target_purchase_table(source: Q1SourceData, plan_a: ScheduleSolution, plan_b: ScheduleSolution) -> pd.DataFrame:
    targets = ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"]
    rows = []
    for interval in targets:
        try:
            idx = source.interval.index(interval)
        except ValueError as exc:
            raise ValueError(f"target interval not found: {interval}") from exc
        rows.append(
            {
                "interval": interval,
                "index": idx + 1,
                "grid_A_kwh": float(plan_a.grid_kwh[idx]),
                "grid_B_kwh": float(plan_b.grid_kwh[idx]),
            }
        )
    return pd.DataFrame(rows)


def _assert_close(name: str, actual: float, expected: float, tol: float = READBACK_TOL) -> None:
    if abs(float(actual) - float(expected)) > tol:
        raise AssertionError(f"{name} mismatch: {actual} != {expected}")


def verify_audit_workbook(path: Path, schedule: pd.DataFrame, summary: pd.DataFrame) -> None:
    read_schedule = pd.read_excel(path, sheet_name="schedule")
    read_summary = pd.read_excel(path, sheet_name="summary")
    if len(read_schedule) != len(schedule):
        raise AssertionError("schedule sheet row count mismatch")
    if len(read_summary) != len(summary):
        raise AssertionError("summary sheet row count mismatch")
    for column in ["grid_A_kwh", "soc_A_kwh", "grid_B_kwh", "soc_B_kwh"]:
        diff = np.nanmax(np.abs(read_schedule[column].to_numpy(dtype=float) - schedule[column].to_numpy(dtype=float)))
        if float(diff) > 1.0e-8:
            raise AssertionError(f"workbook schedule column {column} mismatch: {diff}")


def write_template_result1(path: Path, plan_a: ScheduleSolution) -> Path:
    if not TEMPLATE_RESULT_1.exists():
        raise FileNotFoundError(f"template result1.xlsx not found: {TEMPLATE_RESULT_1}")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(TEMPLATE_RESULT_1)
    purchase_ws = workbook["计划购电量"]
    storage_ws = workbook["充放电量"]

    # The official template labels purchase rows as 0:10-0:20 through 0:00+1-0:10+1.
    # We preserve those labels and fill Plan A values by row order from the physical
    # intervals 00:00-00:10 through 23:50-24:00.
    for i, value in enumerate(plan_a.grid_kwh, start=2):
        purchase_ws.cell(row=i, column=2).value = _display_float(value)

    four_hour = four_hour_table(plan_a)
    for row_index, row in enumerate(four_hour.itertuples(index=False), start=2):
        storage_ws.cell(row=row_index, column=2).value = _display_float(row.charge_A_kwh)
        storage_ws.cell(row=row_index, column=3).value = _display_float(row.discharge_A_kwh)

    storage_ws.cell(row=2, column=5).value = float(SOC_INITIAL_KWH)
    storage_ws.cell(row=3, column=5).value = float(SOC_TERMINAL_KWH)
    workbook.save(path)
    return path


def verify_template_result1(path: Path, plan_a: ScheduleSolution) -> None:
    workbook = load_workbook(path, data_only=True)
    purchase_ws = workbook["计划购电量"]
    storage_ws = workbook["充放电量"]
    if purchase_ws.max_row < INTERVAL_COUNT + 1:
        raise AssertionError("template purchase sheet does not contain 144 data rows")
    for i, expected in enumerate(plan_a.grid_kwh, start=2):
        _assert_close(f"template purchase row {i}", purchase_ws.cell(i, 2).value, expected, tol=1.0e-8)

    four_hour = four_hour_table(plan_a)
    for offset, row in enumerate(four_hour.itertuples(index=False), start=2):
        _assert_close(f"template charge row {offset}", storage_ws.cell(offset, 2).value, row.charge_A_kwh, tol=1.0e-8)
        _assert_close(f"template discharge row {offset}", storage_ws.cell(offset, 3).value, row.discharge_A_kwh, tol=1.0e-8)
    _assert_close("template 0:00 SOC", storage_ws.cell(2, 5).value, SOC_INITIAL_KWH)
    _assert_close("template 24:00 SOC", storage_ws.cell(3, 5).value, SOC_TERMINAL_KWH)


def write_template_mapping_note(path: Path) -> Path:
    text = """# result1.xlsx Mapping Note\n\nSubmitted schedule: Plan A.\n\nThe official template purchase sheet labels rows from `0:10-0:20` through `0:00+1-0:10+1`. The model uses attachment endpoint semantics and physical intervals from `00:00-00:10` through `23:50-24:00`. To preserve the original template labels, `output/question_1/result1.xlsx` fills the 144 Plan A purchase values by row order. The same Plan A schedule is used for four-hour charge/discharge totals. The 0:00 and 24:00 SOC cells are filled explicitly as 6000 kWh.\n"""
    path.write_text(text, encoding="utf-8")
    return path


def export_all(
    outdir: Path,
    source: Q1SourceData,
    interval_data: IntervalData,
    baseline: BaselineSolution,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
    checks: dict[str, SolutionCheck],
) -> dict[str, Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    schedule = schedule_table(source, interval_data, plan_a, plan_b)
    summary = summary_table(baseline, plan_a, plan_b, checks)
    four_hour = four_hour_table(plan_a, plan_b)
    target_purchase = target_purchase_table(source, plan_a, plan_b)

    paths = {
        "schedule_csv": outdir / "q1_schedule.csv",
        "summary_csv": outdir / "q1_summary.csv",
        "four_hour_csv": outdir / "q1_four_hour_summary.csv",
        "target_purchase_csv": outdir / "q1_target_purchase.csv",
        "workbook": outdir / "q1_results.xlsx",
        "template_result1": outdir / "result1.xlsx",
        "template_mapping_note": outdir / "result1_mapping_note.md",
    }

    schedule.to_csv(paths["schedule_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")
    summary.to_csv(paths["summary_csv"], index=False, encoding="utf-8-sig")
    four_hour.to_csv(paths["four_hour_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")
    target_purchase.to_csv(paths["target_purchase_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")

    with pd.ExcelWriter(paths["workbook"], engine="xlsxwriter") as writer:
        schedule.to_excel(writer, sheet_name="schedule", index=False)
        summary.to_excel(writer, sheet_name="summary", index=False)
        four_hour.to_excel(writer, sheet_name="four_hour", index=False)
        target_purchase.to_excel(writer, sheet_name="target_purchase", index=False)

    verify_audit_workbook(paths["workbook"], schedule, summary)
    write_template_result1(paths["template_result1"], plan_a)
    verify_template_result1(paths["template_result1"], plan_a)
    write_template_mapping_note(paths["template_mapping_note"])

    return paths
