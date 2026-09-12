"""Isolated Question 2 experiment for conflict-driven incremental MILP.

This module deliberately reads only ``raw/`` and has no dependency on the
received or formal Question 2 scripts.  It compares solver formulations; it is
not a contest result generator.
"""

from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime, timedelta
from pathlib import Path
from time import perf_counter
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import coo_matrix, csr_matrix


ROOT_DIR = Path(__file__).resolve().parents[1]
RAW_ATTACHMENT_DIR = ROOT_DIR / "raw" / "附件"
SLOTS_PER_DAY = 144
DT_HOURS = 1.0 / 6.0


@dataclass(frozen=True)
class Config:
    eta_c: float = 0.9
    eta_d: float = 0.9
    soc_min_kwh: float = 1200.0
    soc_max_kwh: float = 10800.0
    power_max_kw: float = 5000.0
    initial_soc_kwh: float = 6000.0
    reference_soc_kwh: float = 6000.0
    terminal_penalty: float = 0.5
    emergency_multiplier: float = 5.0
    forecast_window_days: int = 7
    scenario_history_days: int = 28
    scenario_count: int = 20
    random_seed: int = 20260911
    residual_scale: float = 1.0
    mip_time_limit_seconds: float = 60.0
    mip_relative_gap: float = 1e-4
    energy_tolerance_kwh: float = 2e-5

    @property
    def interval_limit_kwh(self) -> float:
        return self.power_max_kw * DT_HOURS

    def validate(self) -> "Config":
        values = np.asarray(
            [
                self.eta_c,
                self.eta_d,
                self.soc_min_kwh,
                self.soc_max_kwh,
                self.power_max_kw,
                self.initial_soc_kwh,
                self.reference_soc_kwh,
                self.terminal_penalty,
                self.emergency_multiplier,
                self.forecast_window_days,
                self.scenario_history_days,
                self.scenario_count,
                self.random_seed,
                self.residual_scale,
                self.mip_time_limit_seconds,
                self.mip_relative_gap,
                self.energy_tolerance_kwh,
            ],
            dtype=float,
        )
        if not np.isfinite(values).all():
            raise ValueError("All configuration values must be finite.")
        if not (0 < self.eta_c <= 1 and 0 < self.eta_d <= 1):
            raise ValueError("Charge and discharge efficiencies must be in (0, 1].")
        if not (0 <= self.soc_min_kwh < self.soc_max_kwh):
            raise ValueError("Invalid SOC bounds.")
        if not (
            self.soc_min_kwh
            <= self.initial_soc_kwh
            <= self.soc_max_kwh
        ):
            raise ValueError("Initial SOC is outside the configured bounds.")
        if not (
            self.soc_min_kwh
            <= self.reference_soc_kwh
            <= self.soc_max_kwh
        ):
            raise ValueError("Reference SOC is outside the configured bounds.")
        if self.power_max_kw <= 0 or self.emergency_multiplier <= 1:
            raise ValueError("Power limit and emergency multiplier are invalid.")
        if min(
            self.forecast_window_days,
            self.scenario_history_days,
            self.scenario_count,
        ) < 1:
            raise ValueError("Forecast and scenario window sizes must be positive.")
        if self.mip_time_limit_seconds <= 0 or self.mip_relative_gap < 0:
            raise ValueError("Invalid MILP stopping settings.")
        return self


@dataclass(frozen=True)
class InputData:
    price: np.ndarray
    load_kwh: np.ndarray
    pv_kwh: np.ndarray
    dates: tuple[datetime, ...]


@dataclass(frozen=True)
class ForecastData:
    load_hat_kwh: np.ndarray
    pv_hat_kwh: np.ndarray
    load_error_kwh: np.ndarray
    pv_error_kwh: np.ndarray
    window_days: int


@dataclass(frozen=True)
class ModelLayout:
    scenario_count: int
    interval_count: int
    q: np.ndarray
    soc: np.ndarray
    emergency: np.ndarray
    terminal_shortfall: np.ndarray
    binary_by_pair: dict[tuple[int, int], int]
    variable_count: int


@dataclass(frozen=True)
class BuiltModel:
    objective: np.ndarray
    lower_bounds: np.ndarray
    upper_bounds: np.ndarray
    constraints: csr_matrix
    constraint_upper_bounds: np.ndarray
    integrality: np.ndarray
    layout: ModelLayout
    build_seconds: float


@dataclass(frozen=True)
class Solution:
    method: str
    success: bool
    status_code: int
    message: str
    objective: float
    dual_bound: float
    relative_gap: float
    build_seconds: float
    solve_seconds: float
    plan_kwh: np.ndarray
    state_kwh: np.ndarray
    emergency_kwh: np.ndarray
    charge_kwh: np.ndarray
    discharge_kwh: np.ndarray
    conflicts: tuple[tuple[int, int], ...]
    active_binary_count: int
    node_count: int
    max_constraint_violation: float


@dataclass(frozen=True)
class IncrementalResult:
    relaxation: Solution
    solution: Solution
    iterations: int
    binary_history: tuple[int, ...]
    total_build_seconds: float
    total_solve_seconds: float


def _read_power_sheet(sheet) -> tuple[np.ndarray, tuple[datetime, ...]]:
    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) != 366 or len(rows[0]) != 145:
        raise ValueError(
            f"{sheet.title} must contain one header row and 365x144 values."
        )
    dates = tuple(row[0] for row in rows[1:])
    if not all(isinstance(value, datetime) for value in dates):
        raise ValueError(f"{sheet.title} contains a non-datetime date value.")
    power_kw = np.asarray([row[1:] for row in rows[1:]], dtype=float)
    return power_kw * DT_HOURS, dates


def load_raw_inputs(raw_attachment_dir: Path = RAW_ATTACHMENT_DIR) -> InputData:
    """Load Question 2 inputs from the immutable early raw-data copy."""

    raw_attachment_dir = Path(raw_attachment_dir).resolve()
    price_path = raw_attachment_dir / "附件1.xlsx"
    actual_path = raw_attachment_dir / "附件2.xlsx"
    if not price_path.is_file() or not actual_path.is_file():
        raise FileNotFoundError(
            f"Expected raw workbooks under {raw_attachment_dir}."
        )

    price_workbook = load_workbook(price_path, data_only=True, read_only=True)
    try:
        price_rows = list(
            price_workbook.worksheets[0].iter_rows(values_only=True)
        )
    finally:
        price_workbook.close()
    if len(price_rows) != 145:
        raise ValueError("Attachment 1 must contain one header and 144 price rows.")
    price = np.asarray([row[1] for row in price_rows[1:]], dtype=float)

    actual_workbook = load_workbook(actual_path, data_only=True, read_only=True)
    try:
        load_sheet = next(
            (
                sheet
                for sheet in actual_workbook.worksheets
                if "负荷" in sheet.title or "负载" in sheet.title
            ),
            None,
        )
        pv_sheet = next(
            (sheet for sheet in actual_workbook.worksheets if "光伏" in sheet.title),
            None,
        )
        if load_sheet is None or pv_sheet is None:
            raise ValueError("Attachment 2 is missing load or photovoltaic sheets.")
        load_kwh, dates = _read_power_sheet(load_sheet)
        pv_kwh, pv_dates = _read_power_sheet(pv_sheet)
    finally:
        actual_workbook.close()

    if dates != pv_dates:
        raise ValueError("Load and photovoltaic dates do not match.")
    expected_dates = tuple(
        datetime(2025, 1, 1) + timedelta(days=day) for day in range(365)
    )
    if dates != expected_dates:
        raise ValueError("Attachment 2 must cover every day of 2025 in order.")
    for name, values in (
        ("price", price),
        ("load", load_kwh),
        ("photovoltaic", pv_kwh),
    ):
        if not np.isfinite(values).all() or np.any(values < 0):
            raise ValueError(f"{name} contains invalid values.")
    if price.shape != (SLOTS_PER_DAY,) or np.any(price <= 0):
        raise ValueError("Price must contain 144 strictly positive values.")
    return InputData(price, load_kwh, pv_kwh, dates)


def prepare_causal_forecast(
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    window_days: int,
) -> ForecastData:
    load_kwh = np.asarray(load_kwh, dtype=float)
    pv_kwh = np.asarray(pv_kwh, dtype=float)
    if load_kwh.shape != pv_kwh.shape or load_kwh.ndim != 2:
        raise ValueError("Load and photovoltaic arrays must have equal 2-D shapes.")
    if window_days < 1:
        raise ValueError("Forecast window must be positive.")
    load_hat = np.zeros_like(load_kwh)
    pv_hat = np.zeros_like(pv_kwh)
    for day in range(1, len(load_kwh)):
        start = max(0, day - window_days)
        load_hat[day] = load_kwh[start:day].mean(axis=0)
        pv_hat[day] = pv_kwh[start:day].mean(axis=0)
    return ForecastData(
        load_hat,
        pv_hat,
        load_kwh - load_hat,
        pv_kwh - pv_hat,
        window_days,
    )


def make_joint_scenarios(
    day: int,
    forecast: ForecastData,
    load_kwh: np.ndarray,
    pv_kwh: np.ndarray,
    config: Config,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    if not 0 <= day < len(load_kwh):
        raise IndexError("Target day is outside the input range.")
    start = max(
        forecast.window_days,
        day - config.scenario_history_days,
    )
    pool = np.arange(start, day, dtype=int)
    if pool.size == 0:
        sources = np.asarray([-1], dtype=int)
        scenario_load = forecast.load_hat_kwh[day][None, :].copy()
        scenario_pv = forecast.pv_hat_kwh[day][None, :].copy()
    else:
        generator = np.random.default_rng(config.random_seed + day)
        count = min(config.scenario_count, pool.size)
        sources = np.sort(generator.choice(pool, size=count, replace=False))
        scenario_load = np.maximum(
            0.0,
            forecast.load_hat_kwh[day][None, :]
            + config.residual_scale * forecast.load_error_kwh[sources],
        )
        scenario_pv = np.maximum(
            0.0,
            forecast.pv_hat_kwh[day][None, :]
            + config.residual_scale * forecast.pv_error_kwh[sources],
        )
    weights = np.full(len(sources), 1.0 / len(sources))
    if np.any(sources >= day):
        raise AssertionError("A scenario source is not earlier than the target day.")
    return scenario_load, scenario_pv, weights, sources


def _validate_model_inputs(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    config.validate()
    load = np.atleast_2d(np.asarray(load_scenarios, dtype=float))
    pv = np.atleast_2d(np.asarray(pv_scenarios, dtype=float))
    probability = np.asarray(weights, dtype=float)
    price = np.asarray(price, dtype=float)
    if load.shape != pv.shape or load.shape[1] != len(price):
        raise ValueError("Scenario and price dimensions do not match.")
    if len(probability) != load.shape[0] or np.any(probability < 0):
        raise ValueError("Scenario probabilities are invalid.")
    if not np.isclose(probability.sum(), 1.0):
        raise ValueError("Scenario probabilities must sum to one.")
    if not np.isfinite(load).all() or not np.isfinite(pv).all():
        raise ValueError("Scenario data must be finite.")
    if np.any(load < 0) or np.any(pv < 0) or np.any(price <= 0):
        raise ValueError("Scenario energy must be nonnegative and price positive.")
    if not config.soc_min_kwh <= initial_soc_kwh <= config.soc_max_kwh:
        raise ValueError("Initial SOC is outside the configured bounds.")
    return load, pv, probability, price


def build_model(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
    active_binary_pairs: Iterable[tuple[int, int]] = (),
    terminal_free: bool = False,
) -> BuiltModel:
    """Build the common relaxation plus binaries for selected scenario slots."""

    started = perf_counter()
    load, pv, probability, price = _validate_model_inputs(
        load_scenarios,
        pv_scenarios,
        weights,
        price,
        initial_soc_kwh,
        config,
    )
    scenario_count, interval_count = load.shape
    pair_count = scenario_count * interval_count
    q = np.arange(interval_count, dtype=int)
    soc = interval_count + np.arange(pair_count).reshape(
        scenario_count, interval_count
    )
    emergency = interval_count + pair_count + np.arange(pair_count).reshape(
        scenario_count, interval_count
    )
    terminal_shortfall = (
        interval_count + 2 * pair_count + np.arange(scenario_count)
    )
    continuous_count = interval_count + 2 * pair_count + scenario_count

    normalized_pairs = tuple(sorted(set(active_binary_pairs)))
    for scenario, interval in normalized_pairs:
        if not (
            0 <= scenario < scenario_count
            and 0 <= interval < interval_count
        ):
            raise IndexError(f"Invalid binary pair {(scenario, interval)}.")
    binary_by_pair = {
        pair: continuous_count + offset
        for offset, pair in enumerate(normalized_pairs)
    }
    variable_count = continuous_count + len(normalized_pairs)
    layout = ModelLayout(
        scenario_count,
        interval_count,
        q,
        soc,
        emergency,
        terminal_shortfall,
        binary_by_pair,
        variable_count,
    )

    objective = np.zeros(variable_count)
    objective[q] = price
    objective[emergency] = (
        probability[:, None]
        * config.emergency_multiplier
        * price[None, :]
    )
    if not terminal_free:
        objective[terminal_shortfall] = probability * config.terminal_penalty

    lower = np.zeros(variable_count)
    upper = np.full(variable_count, np.inf)
    lower[soc] = config.soc_min_kwh
    upper[soc] = config.soc_max_kwh
    upper[emergency] = load
    for index in binary_by_pair.values():
        upper[index] = 1.0

    rows: list[int] = []
    columns: list[int] = []
    values: list[float] = []
    right_hand_side: list[float] = []

    def add_constraint(
        coefficients: Iterable[tuple[int, float]],
        bound: float,
    ) -> None:
        row = len(right_hand_side)
        right_hand_side.append(float(bound))
        for column, coefficient in coefficients:
            rows.append(row)
            columns.append(int(column))
            values.append(float(coefficient))

    def state_delta(
        scenario: int,
        interval: int,
        scale: float,
    ) -> list[tuple[int, float]]:
        terms = [(soc[scenario, interval], scale)]
        if interval > 0:
            terms.append((soc[scenario, interval - 1], -scale))
        return terms

    charge_soc_limit = config.eta_c * config.interval_limit_kwh
    discharge_soc_limit = (
        np.minimum(config.interval_limit_kwh, load) / config.eta_d
    )
    for scenario in range(scenario_count):
        for interval in range(interval_count):
            for slope in (1.0 / config.eta_c, config.eta_d):
                add_constraint(
                    [
                        (q[interval], -1.0),
                        (emergency[scenario, interval], -1.0),
                    ]
                    + state_delta(scenario, interval, slope),
                    -(load[scenario, interval] - pv[scenario, interval])
                    + (slope * initial_soc_kwh if interval == 0 else 0.0),
                )
            add_constraint(
                state_delta(scenario, interval, 1.0),
                charge_soc_limit
                + (initial_soc_kwh if interval == 0 else 0.0),
            )
            add_constraint(
                state_delta(scenario, interval, -1.0),
                discharge_soc_limit[scenario, interval]
                - (initial_soc_kwh if interval == 0 else 0.0),
            )
        add_constraint(
            [
                (soc[scenario, -1], -1.0),
                (terminal_shortfall[scenario], -1.0),
            ],
            -config.reference_soc_kwh,
        )

    for (scenario, interval), binary in binary_by_pair.items():
        add_constraint(
            state_delta(scenario, interval, 1.0)
            + [(binary, -charge_soc_limit)],
            initial_soc_kwh if interval == 0 else 0.0,
        )
        interval_discharge_limit = discharge_soc_limit[scenario, interval]
        add_constraint(
            state_delta(scenario, interval, -1.0)
            + [(binary, interval_discharge_limit)],
            interval_discharge_limit
            - (initial_soc_kwh if interval == 0 else 0.0),
        )
        add_constraint(
            [
                (emergency[scenario, interval], 1.0),
                (binary, load[scenario, interval]),
            ],
            load[scenario, interval],
        )

    constraint_matrix = coo_matrix(
        (values, (rows, columns)),
        shape=(len(right_hand_side), variable_count),
    ).tocsr()
    integrality = np.zeros(variable_count, dtype=int)
    if normalized_pairs:
        integrality[continuous_count:] = 1
    return BuiltModel(
        objective,
        lower,
        upper,
        constraint_matrix,
        np.asarray(right_hand_side),
        integrality,
        layout,
        perf_counter() - started,
    )


def _decode_solution(
    method: str,
    raw_result,
    model: BuiltModel,
    initial_soc_kwh: float,
    config: Config,
    solve_seconds: float,
) -> Solution:
    if raw_result.x is None:
        raise RuntimeError(f"{method} returned no feasible solution: {raw_result.message}")
    vector = np.asarray(raw_result.x, dtype=float)
    layout = model.layout
    state = np.column_stack(
        [
            np.full(layout.scenario_count, initial_soc_kwh),
            vector[layout.soc],
        ]
    )
    state_change = np.diff(state, axis=1)
    emergency = vector[layout.emergency]
    charge = np.maximum(state_change, 0.0) / config.eta_c
    discharge = config.eta_d * np.maximum(-state_change, 0.0)
    conflicts_array = np.argwhere(
        (charge > config.energy_tolerance_kwh)
        & (emergency > config.energy_tolerance_kwh)
    )
    conflicts = tuple(map(tuple, conflicts_array.tolist()))
    violation = model.constraints @ vector - model.constraint_upper_bounds
    max_violation = max(0.0, float(np.max(violation)))
    if not np.isfinite(vector).all():
        raise AssertionError(f"{method} returned non-finite variables.")
    if np.min(vector[layout.q]) < -config.energy_tolerance_kwh:
        raise AssertionError(f"{method} returned a negative plan.")
    if np.min(emergency) < -config.energy_tolerance_kwh:
        raise AssertionError(f"{method} returned negative emergency energy.")
    if state.min() < config.soc_min_kwh - config.energy_tolerance_kwh:
        raise AssertionError(f"{method} violated the lower SOC bound.")
    if state.max() > config.soc_max_kwh + config.energy_tolerance_kwh:
        raise AssertionError(f"{method} violated the upper SOC bound.")
    if charge.max() > config.interval_limit_kwh + config.energy_tolerance_kwh:
        raise AssertionError(f"{method} violated the charge limit.")
    if discharge.max() > config.interval_limit_kwh + config.energy_tolerance_kwh:
        raise AssertionError(f"{method} violated the discharge limit.")
    if max_violation > config.energy_tolerance_kwh:
        raise AssertionError(
            f"{method} constraint violation {max_violation} exceeds tolerance."
        )
    relative_gap = float(getattr(raw_result, "mip_gap", 0.0) or 0.0)
    dual_bound = float(getattr(raw_result, "mip_dual_bound", raw_result.fun))
    node_count = int(getattr(raw_result, "mip_node_count", 0) or 0)
    return Solution(
        method=method,
        success=bool(raw_result.success),
        status_code=int(raw_result.status),
        message=str(raw_result.message),
        objective=float(model.objective @ vector),
        dual_bound=dual_bound,
        relative_gap=relative_gap,
        build_seconds=model.build_seconds,
        solve_seconds=solve_seconds,
        plan_kwh=np.maximum(vector[layout.q], 0.0),
        state_kwh=state,
        emergency_kwh=emergency,
        charge_kwh=charge,
        discharge_kwh=discharge,
        conflicts=conflicts,
        active_binary_count=len(layout.binary_by_pair),
        node_count=node_count,
        max_constraint_violation=max_violation,
    )


def solve_with_binary_pairs(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
    active_binary_pairs: Iterable[tuple[int, int]],
    method: str,
    terminal_free: bool = False,
) -> Solution:
    pairs = tuple(active_binary_pairs)
    model = build_model(
        load_scenarios,
        pv_scenarios,
        weights,
        price,
        initial_soc_kwh,
        config,
        pairs,
        terminal_free,
    )
    started = perf_counter()
    if not pairs:
        raw_result = linprog(
            model.objective,
            A_ub=model.constraints,
            b_ub=model.constraint_upper_bounds,
            bounds=list(zip(model.lower_bounds, model.upper_bounds)),
            method="highs",
        )
    else:
        raw_result = milp(
            model.objective,
            integrality=model.integrality,
            bounds=Bounds(model.lower_bounds, model.upper_bounds),
            constraints=LinearConstraint(
                model.constraints,
                -np.inf,
                model.constraint_upper_bounds,
            ),
            options={
                "time_limit": config.mip_time_limit_seconds,
                "mip_rel_gap": config.mip_relative_gap,
            },
        )
    solve_seconds = perf_counter() - started
    return _decode_solution(
        method,
        raw_result,
        model,
        initial_soc_kwh,
        config,
        solve_seconds,
    )


def solve_relaxation(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
    terminal_free: bool = False,
) -> Solution:
    return solve_with_binary_pairs(
        load_scenarios,
        pv_scenarios,
        weights,
        price,
        initial_soc_kwh,
        config,
        (),
        "LP_RELAXATION",
        terminal_free,
    )


def solve_full_milp(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
    terminal_free: bool = False,
) -> Solution:
    load = np.atleast_2d(load_scenarios)
    pairs = (
        (scenario, interval)
        for scenario in range(load.shape[0])
        for interval in range(load.shape[1])
    )
    result = solve_with_binary_pairs(
        load_scenarios,
        pv_scenarios,
        weights,
        price,
        initial_soc_kwh,
        config,
        pairs,
        "FULL_MILP",
        terminal_free,
    )
    if result.conflicts:
        raise AssertionError("Full MILP returned a charge/emergency conflict.")
    return result


def solve_incremental_milp(
    load_scenarios: np.ndarray,
    pv_scenarios: np.ndarray,
    weights: np.ndarray,
    price: np.ndarray,
    initial_soc_kwh: float,
    config: Config,
    terminal_free: bool = False,
    max_iterations: int = 20,
    activation_scope: str = "scenario",
    trace: bool = False,
) -> IncrementalResult:
    """Activate binaries lazily until the returned solution is full-MILP feasible."""

    if max_iterations < 1:
        raise ValueError("max_iterations must be positive.")
    if activation_scope not in {"pair", "scenario", "interval"}:
        raise ValueError(
            "activation_scope must be 'pair', 'scenario', or 'interval'."
        )
    relaxation = solve_relaxation(
        load_scenarios,
        pv_scenarios,
        weights,
        price,
        initial_soc_kwh,
        config,
        terminal_free,
    )
    total_build = relaxation.build_seconds
    total_solve = relaxation.solve_seconds
    scenario_count, interval_count = np.atleast_2d(load_scenarios).shape

    def expand(
        conflicts: Iterable[tuple[int, int]],
    ) -> set[tuple[int, int]]:
        conflicts = tuple(conflicts)
        if activation_scope == "pair":
            return set(conflicts)
        if activation_scope == "scenario":
            scenarios = {scenario for scenario, _ in conflicts}
            return {
                (scenario, interval)
                for scenario in scenarios
                for interval in range(interval_count)
            }
        intervals = {interval for _, interval in conflicts}
        return {
            (scenario, interval)
            for scenario in range(scenario_count)
            for interval in intervals
        }

    active = expand(relaxation.conflicts)
    binary_history = [0]
    if trace:
        print(
            "[incremental] relaxation "
            f"objective={relaxation.objective:.9g} "
            f"raw_conflicts={len(relaxation.conflicts)} "
            f"initial_active={len(active)} "
            f"solve={relaxation.solve_seconds:.3f}s",
            flush=True,
        )
    if not active:
        certified = replace(relaxation, method="INCREMENTAL_LP_CERTIFIED")
        return IncrementalResult(
            relaxation,
            certified,
            0,
            tuple(binary_history),
            total_build,
            total_solve,
        )

    for iteration in range(1, max_iterations + 1):
        binary_history.append(len(active))
        solution = solve_with_binary_pairs(
            load_scenarios,
            pv_scenarios,
            weights,
            price,
            initial_soc_kwh,
            config,
            active,
            f"INCREMENTAL_MILP_ROUND_{iteration}",
            terminal_free,
        )
        total_build += solution.build_seconds
        total_solve += solution.solve_seconds
        if not solution.success:
            raise RuntimeError(
                "Incremental MILP did not reach its requested stopping condition: "
                f"{solution.message}"
            )
        unresolved = set(solution.conflicts)
        new_conflicts = unresolved - active
        if trace:
            print(
                f"[incremental] round={iteration} active={len(active)} "
                f"unresolved={len(unresolved)} new={len(new_conflicts)} "
                f"objective={solution.objective:.9g} "
                f"gap={solution.relative_gap:.3g} "
                f"solve={solution.solve_seconds:.3f}s",
                flush=True,
            )
        if not unresolved:
            return IncrementalResult(
                relaxation,
                replace(solution, method="INCREMENTAL_MILP_CERTIFIED"),
                iteration,
                tuple(binary_history),
                total_build,
                total_solve,
            )
        if not new_conflicts:
            raise AssertionError(
                "Active binary constraints still contain a conflict above tolerance."
            )
        active.update(expand(new_conflicts))
    raise RuntimeError(
        "Incremental MILP did not remove all conflicts in "
        f"{max_iterations} rounds; binary history={binary_history}, "
        f"last unresolved={len(unresolved)}, active={len(active)}."
    )


def objective_interval(solution: Solution) -> tuple[float, float]:
    """Return the solver's numerical lower and upper objective bounds."""

    lower = min(solution.dual_bound, solution.objective)
    upper = max(solution.dual_bound, solution.objective)
    return lower, upper


def numerical_objective_agreement(
    first: Solution,
    second: Solution,
    absolute_tolerance: float = 1e-5,
) -> bool:
    """Check whether two solver bound intervals overlap within tolerance."""

    first_lower, first_upper = objective_interval(first)
    second_lower, second_upper = objective_interval(second)
    return max(first_lower, second_lower) <= (
        min(first_upper, second_upper) + absolute_tolerance
    )
