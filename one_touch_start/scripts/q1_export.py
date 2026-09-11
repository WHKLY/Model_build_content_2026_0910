"""Question 1 table, workbook, and official-template exports."""
from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
from pathlib import PurePosixPath
import tempfile
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from project_paths import TEMPLATE_RESULT_1
from q1_data import INTERVAL_COUNT, Q1SourceData
from q1_model import BaselineSolution, IntervalData, SOC_INITIAL_KWH, SOC_TERMINAL_KWH, ScheduleSolution
from q1_verify import SolutionCheck

READBACK_TOL = 1.0e-6
DISPLAY_ZERO_TOL = 1.0e-8
SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
DOCUMENT_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
NS = {"main": SPREADSHEET_NS, "rel": DOCUMENT_REL_NS, "pkg": PACKAGE_REL_NS}

TEMPLATE_EDITABLE_CELLS = {
    "计划购电量": {f"B{row}" for row in range(2, INTERVAL_COUNT + 2)},
    "充放电量": {
        *(f"B{row}" for row in range(2, 8)),
        *(f"C{row}" for row in range(2, 8)),
        "E2",
        "E3",
    },
}


def _display_float(value: float) -> float:
    value = float(value)
    return 0.0 if abs(value) < DISPLAY_ZERO_TOL else value


def schedule_table(
    source: Q1SourceData,
    interval_data: IntervalData,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "index": np.arange(1, INTERVAL_COUNT + 1),
            "source_time": source.source_time,
            "interval": source.interval,
            "price_yuan_per_kwh": interval_data.price,
            "load_kw": interval_data.load_kw,
            "pv_forecast_kw": interval_data.pv_kw,
            "load_kwh": interval_data.load_kwh,
            "pv_forecast_kwh": interval_data.pv_kwh,
            "grid_A_kwh": plan_a.grid_kwh,
            "charge_A_kwh": plan_a.charge_kwh,
            "discharge_A_kwh": plan_a.discharge_kwh,
            "curtail_A_kwh": plan_a.curtail_kwh,
            "soc_A_kwh": plan_a.soc_kwh,
            "grid_B_kwh": plan_b.grid_kwh,
            "charge_B_kwh": plan_b.charge_kwh,
            "discharge_B_kwh": plan_b.discharge_kwh,
            "curtail_B_kwh": plan_b.curtail_kwh,
            "soc_B_kwh": plan_b.soc_kwh,
            "diff_grid_B_minus_A_kwh": plan_b.grid_kwh - plan_a.grid_kwh,
            "diff_soc_B_minus_A_kwh": plan_b.soc_kwh - plan_a.soc_kwh,
        }
    )


def summary_table(
    baseline: BaselineSolution,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
    checks: dict[str, SolutionCheck],
    certificate: dict[str, Any],
) -> pd.DataFrame:
    rows = [
        ("no_storage_cost_yuan", baseline.cost),
        ("no_storage_grid_kwh", baseline.total_grid_kwh),
        ("plan_A_cost_yuan", plan_a.cost),
        ("plan_A_grid_kwh", plan_a.total_grid_kwh),
        ("plan_A_charge_kwh", plan_a.total_charge_kwh),
        ("plan_A_discharge_kwh", plan_a.total_discharge_kwh),
        ("plan_A_curtail_kwh", _display_float(plan_a.total_curtail_kwh)),
        ("plan_A_target_index", plan_a.target_index),
        ("plan_A_target_grid_kwh", plan_a.target_grid_kwh),
        ("plan_A_max_balance_residual", checks["plan_a"].max_balance_residual),
        ("plan_A_max_soc_recurrence_residual", checks["plan_a"].max_soc_recurrence_residual),
        ("plan_A_max_lp_ub_violation", checks["plan_a"].max_lp_ub_violation),
        ("plan_A_max_lp_eq_residual", checks["plan_a"].max_lp_eq_residual),
        ("plan_B_cost_yuan", plan_b.cost),
        ("plan_B_grid_kwh", plan_b.total_grid_kwh),
        ("plan_B_charge_kwh", plan_b.total_charge_kwh),
        ("plan_B_discharge_kwh", plan_b.total_discharge_kwh),
        ("plan_B_curtail_kwh", _display_float(plan_b.total_curtail_kwh)),
        ("plan_B_target_index", plan_b.target_index),
        ("plan_B_target_grid_kwh", plan_b.target_grid_kwh),
        ("plan_B_max_balance_residual", checks["plan_b"].max_balance_residual),
        ("plan_B_max_soc_recurrence_residual", checks["plan_b"].max_soc_recurrence_residual),
        ("plan_B_max_lp_ub_violation", checks["plan_b"].max_lp_ub_violation),
        ("plan_B_max_lp_eq_residual", checks["plan_b"].max_lp_eq_residual),
        ("cost_saving_A_vs_no_storage_yuan", baseline.cost - plan_a.cost),
        ("cost_saving_B_vs_no_storage_yuan", baseline.cost - plan_b.cost),
        ("max_abs_grid_difference_B_minus_A_kwh", float(np.max(np.abs(plan_b.grid_kwh - plan_a.grid_kwh)))),
        ("certificate_upper_bound_exact", certificate["upper_bound"]),
        ("certificate_lower_bound_exact", certificate["lower_bound"]),
        ("certificate_gap_exact", certificate["gap"]),
        ("certificate_exact_optimal", certificate["exact_optimal"]),
        ("model_note", "same optimal purchase cost, different interval-level purchase schedules"),
    ]
    return pd.DataFrame(rows, columns=["metric", "value"])


def four_hour_table(plan_a: ScheduleSolution, plan_b: ScheduleSolution | None = None) -> pd.DataFrame:
    labels = ["00:00-04:00", "04:00-08:00", "08:00-12:00", "12:00-16:00", "16:00-20:00", "20:00-24:00"]
    rows = []
    for k, label in enumerate(labels):
        idx = slice(k * 24, (k + 1) * 24)
        row = {
            "interval": label,
            "charge_A_kwh": float(np.sum(plan_a.charge_kwh[idx])),
            "discharge_A_kwh": float(np.sum(plan_a.discharge_kwh[idx])),
            "grid_A_kwh": float(np.sum(plan_a.grid_kwh[idx])),
        }
        if plan_b is not None:
            row.update(
                {
                    "charge_B_kwh": float(np.sum(plan_b.charge_kwh[idx])),
                    "discharge_B_kwh": float(np.sum(plan_b.discharge_kwh[idx])),
                    "grid_B_kwh": float(np.sum(plan_b.grid_kwh[idx])),
                }
            )
        rows.append(row)
    return pd.DataFrame(rows)


def target_purchase_table(source: Q1SourceData, plan_a: ScheduleSolution, plan_b: ScheduleSolution) -> pd.DataFrame:
    targets = ["10:00-10:10", "12:00-12:10", "14:00-14:10", "16:00-16:10", "18:00-18:10", "20:00-20:10"]
    rows = []
    for interval in targets:
        try:
            idx = source.interval.index(interval)
        except ValueError as exc:
            raise ValueError(f"target interval not found: {interval}") from exc
        rows.append(
            {
                "interval": interval,
                "index": idx + 1,
                "grid_A_kwh": float(plan_a.grid_kwh[idx]),
                "grid_B_kwh": float(plan_b.grid_kwh[idx]),
            }
        )
    return pd.DataFrame(rows)


def _assert_close(name: str, actual: float, expected: float, tol: float = READBACK_TOL) -> None:
    if abs(float(actual) - float(expected)) > tol:
        raise AssertionError(f"{name} mismatch: {actual} != {expected}")


def certificate_table(certificate: dict[str, Any], verification: dict[str, Any]) -> pd.DataFrame:
    rows = [
        ("certificate_format", certificate["format"]),
        ("input_file", certificate["input_file"]),
        ("input_sha256", certificate["input_sha256"]),
        ("primal_upper_bound_exact", certificate["upper_bound"]),
        ("dual_lower_bound_exact", certificate["lower_bound"]),
        ("exact_gap", certificate["gap"]),
        ("exact_global_optimal", certificate["exact_optimal"]),
        ("independent_verifier_source_checked", verification["verified_source"]),
        ("verification_note", "See certificate_q1.json and run scripts/q1_certificate_verify.py."),
    ]
    return pd.DataFrame(rows, columns=["check", "value"])


def verify_audit_workbook(
    path: Path,
    schedule: pd.DataFrame,
    summary: pd.DataFrame,
    certificate_summary: pd.DataFrame,
) -> None:
    read_schedule = pd.read_excel(path, sheet_name="schedule")
    read_summary = pd.read_excel(path, sheet_name="summary")
    read_certificate = pd.read_excel(path, sheet_name="certificate")
    if len(read_schedule) != len(schedule):
        raise AssertionError("schedule sheet row count mismatch")
    if len(read_summary) != len(summary):
        raise AssertionError("summary sheet row count mismatch")
    if len(read_certificate) != len(certificate_summary):
        raise AssertionError("certificate sheet row count mismatch")
    certificate_values = dict(zip(read_certificate["check"], read_certificate["value"]))
    if str(certificate_values.get("exact_gap")) != "0":
        raise AssertionError("certificate sheet exact gap is not zero")
    if not bool(certificate_values.get("exact_global_optimal")):
        raise AssertionError("certificate sheet does not report exact global optimality")
    for column in ["grid_A_kwh", "soc_A_kwh", "grid_B_kwh", "soc_B_kwh"]:
        diff = np.nanmax(np.abs(read_schedule[column].to_numpy(dtype=float) - schedule[column].to_numpy(dtype=float)))
        if float(diff) > 1.0e-8:
            raise AssertionError(f"workbook schedule column {column} mismatch: {diff}")


def _validated_template_path() -> Path:
    if not TEMPLATE_RESULT_1.exists():
        raise FileNotFoundError(f"template result1.xlsx not found: {TEMPLATE_RESULT_1}")
    return TEMPLATE_RESULT_1


def _register_xml_namespaces(xml_bytes: bytes) -> None:
    for _, namespace in ET.iterparse(BytesIO(xml_bytes), events=("start-ns",)):
        prefix, uri = namespace
        ET.register_namespace(prefix or "", uri)


def _worksheet_parts(archive: ZipFile) -> dict[str, str]:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    relationships_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationships = {
        node.attrib["Id"]: node.attrib["Target"]
        for node in relationships_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    parts: dict[str, str] = {}
    for sheet in workbook_root.findall("main:sheets/main:sheet", NS):
        relationship_id = sheet.attrib[f"{{{DOCUMENT_REL_NS}}}id"]
        target = relationships[relationship_id]
        if target.startswith("/"):
            part = target.lstrip("/")
        else:
            part = str(PurePosixPath("xl") / PurePosixPath(target))
        parts[sheet.attrib["name"]] = part
    return parts


def _column_number(cell_reference: str) -> int:
    value = 0
    for character in cell_reference:
        if not character.isalpha():
            break
        value = value * 26 + ord(character.upper()) - ord("A") + 1
    return value


def _set_numeric_cell(worksheet_root: ET.Element, cell_reference: str, value: float) -> None:
    row_number = int("".join(character for character in cell_reference if character.isdigit()))
    sheet_data = worksheet_root.find("main:sheetData", NS)
    if sheet_data is None:
        raise AssertionError("template worksheet has no sheetData element")
    row = sheet_data.find(f"main:row[@r='{row_number}']", NS)
    if row is None:
        raise AssertionError(f"template worksheet has no row {row_number}")

    cell = row.find(f"main:c[@r='{cell_reference}']", NS)
    if cell is None:
        cell = ET.Element(f"{{{SPREADSHEET_NS}}}c", {"r": cell_reference})
        target_column = _column_number(cell_reference)
        insert_at = len(row)
        for index, existing in enumerate(row.findall("main:c", NS)):
            if _column_number(existing.attrib["r"]) > target_column:
                insert_at = index
                break
        row.insert(insert_at, cell)

    cell.attrib.pop("t", None)
    for child in list(cell):
        if child.tag in {
            f"{{{SPREADSHEET_NS}}}f",
            f"{{{SPREADSHEET_NS}}}is",
            f"{{{SPREADSHEET_NS}}}v",
        }:
            cell.remove(child)
    value_node = ET.SubElement(cell, f"{{{SPREADSHEET_NS}}}v")
    value_node.text = repr(_display_float(value))


def _patch_worksheet(xml_bytes: bytes, values: dict[str, float]) -> bytes:
    _register_xml_namespaces(xml_bytes)
    root = ET.fromstring(xml_bytes)
    for cell_reference, value in values.items():
        _set_numeric_cell(root, cell_reference, value)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _template_values(plan_a: ScheduleSolution) -> dict[str, dict[str, float]]:
    purchase_values = {
        f"B{row}": float(value)
        for row, value in enumerate(plan_a.grid_kwh, start=2)
    }
    four_hour = four_hour_table(plan_a)
    storage_values: dict[str, float] = {}
    for row_index, row in enumerate(four_hour.itertuples(index=False), start=2):
        storage_values[f"B{row_index}"] = float(row.charge_A_kwh)
        storage_values[f"C{row_index}"] = float(row.discharge_A_kwh)
    storage_values["E2"] = float(SOC_INITIAL_KWH)
    storage_values["E3"] = float(SOC_TERMINAL_KWH)
    return {"计划购电量": purchase_values, "充放电量": storage_values}


def write_template_result1(path: Path, source: Q1SourceData, plan_a: ScheduleSolution) -> Path:
    del source  # Row order is fixed by the official template; labels remain untouched.
    template_path = _validated_template_path()
    path.parent.mkdir(parents=True, exist_ok=True)
    values_by_sheet = _template_values(plan_a)

    temp_file = tempfile.NamedTemporaryFile(
        prefix=f"{path.stem}-",
        suffix=".xlsx",
        dir=path.parent,
        delete=False,
    )
    temp_path = Path(temp_file.name)
    temp_file.close()
    try:
        with ZipFile(template_path, "r") as source_archive, ZipFile(temp_path, "w") as output_archive:
            worksheet_parts = _worksheet_parts(source_archive)
            replacements = {
                worksheet_parts[sheet_name]: values
                for sheet_name, values in values_by_sheet.items()
            }
            output_archive.comment = source_archive.comment
            for item in source_archive.infolist():
                payload = source_archive.read(item.filename)
                if item.filename in replacements:
                    payload = _patch_worksheet(payload, replacements[item.filename])
                output_archive.writestr(item, payload)
        os.replace(temp_path, path)
    finally:
        if temp_path.exists():
            temp_path.unlink()
    return path


def _xml_signature(element: ET.Element) -> tuple[Any, ...]:
    text = element.text.strip() if element.text and element.text.strip() else None
    return (
        element.tag,
        tuple(sorted(element.attrib.items())),
        text,
        tuple(_xml_signature(child) for child in element),
    )


def _worksheet_signature_without_editable_cells(xml_bytes: bytes, editable_cells: set[str]) -> tuple[Any, ...]:
    root = ET.fromstring(xml_bytes)
    for row in root.findall("main:sheetData/main:row", NS):
        for cell in list(row.findall("main:c", NS)):
            if cell.attrib.get("r") in editable_cells:
                row.remove(cell)
    return _xml_signature(root)


def _verify_template_package(path: Path, template_path: Path) -> None:
    with ZipFile(template_path, "r") as template_archive, ZipFile(path, "r") as output_archive:
        template_parts = _worksheet_parts(template_archive)
        output_parts = _worksheet_parts(output_archive)
        if template_parts != output_parts:
            raise AssertionError("result1.xlsx worksheet package mapping differs from the template")
        if template_archive.namelist() != output_archive.namelist():
            raise AssertionError("result1.xlsx package members differ from the template")

        editable_parts = set(template_parts.values())
        for item in template_archive.infolist():
            if item.filename not in editable_parts:
                if template_archive.read(item.filename) != output_archive.read(item.filename):
                    raise AssertionError(f"unapproved template package part changed: {item.filename}")

        for sheet_name, editable_cells in TEMPLATE_EDITABLE_CELLS.items():
            part = template_parts[sheet_name]
            template_signature = _worksheet_signature_without_editable_cells(
                template_archive.read(part), editable_cells
            )
            output_signature = _worksheet_signature_without_editable_cells(
                output_archive.read(part), editable_cells
            )
            if template_signature != output_signature:
                raise AssertionError(f"unapproved worksheet structure changed: {sheet_name}")


def verify_template_result1(path: Path, source: Q1SourceData, plan_a: ScheduleSolution) -> None:
    del source
    template_path = _validated_template_path()
    _verify_template_package(path, template_path)
    template_workbook = load_workbook(template_path, data_only=False)
    workbook = load_workbook(path, data_only=True)
    if workbook.sheetnames != template_workbook.sheetnames:
        raise AssertionError("result1.xlsx worksheet names or order differ from the template")
    purchase_ws = workbook["计划购电量"]
    storage_ws = workbook["充放电量"]
    template_purchase_ws = template_workbook["计划购电量"]
    template_storage_ws = template_workbook["充放电量"]
    if purchase_ws.max_row < INTERVAL_COUNT + 1:
        raise AssertionError("template purchase sheet does not contain 144 data rows")
    for i, expected in enumerate(plan_a.grid_kwh, start=2):
        if purchase_ws.cell(i, 1).value != template_purchase_ws.cell(i, 1).value:
            raise AssertionError(f"template purchase row {i} label changed")
        if purchase_ws.cell(i, 2).style_id != template_purchase_ws.cell(i, 2).style_id:
            raise AssertionError(f"template purchase row {i} style changed")
        _assert_close(f"template purchase row {i}", purchase_ws.cell(i, 2).value, expected, tol=1.0e-8)

    four_hour = four_hour_table(plan_a)
    for offset, row in enumerate(four_hour.itertuples(index=False), start=2):
        if storage_ws.cell(offset, 1).value != template_storage_ws.cell(offset, 1).value:
            raise AssertionError(f"template storage row {offset} label changed")
        _assert_close(f"template charge row {offset}", storage_ws.cell(offset, 2).value, row.charge_A_kwh, tol=1.0e-8)
        _assert_close(f"template discharge row {offset}", storage_ws.cell(offset, 3).value, row.discharge_A_kwh, tol=1.0e-8)
    for coordinate in ("D2", "D3"):
        if storage_ws[coordinate].value != template_storage_ws[coordinate].value:
            raise AssertionError(f"template storage cell {coordinate} changed")
    _assert_close("template 0:00 SOC", storage_ws.cell(2, 5).value, SOC_INITIAL_KWH)
    _assert_close("template 24:00 SOC", storage_ws.cell(3, 5).value, SOC_TERMINAL_KWH)


def write_template_mapping_note(path: Path) -> Path:
    text = """# result1.xlsx Mapping Note\n\nSubmitted schedule: Plan A.\n\nThe generated `output/result1.xlsx` strictly preserves the official template at `assets/附件5/result1.xlsx`, including its original purchase-row labels, formatting, worksheets, shared strings, and printer settings. Only the designated blank numeric result cells are populated.\n\nThe model interprets Attachment 1 labels as interval endpoints and therefore covers physical intervals `00:00-00:10` through `23:50-24:00`. Those physical labels are retained in the audit CSV and `q1_results.xlsx`; they are not written over the official template. Purchase values are written in the template's existing row order. Four-hour charge/discharge totals and the explicit 0:00/24:00 SOC values use the same Plan A schedule.\n"""
    path.write_text(text, encoding="utf-8")
    return path


def export_all(
    outdir: Path,
    source: Q1SourceData,
    interval_data: IntervalData,
    baseline: BaselineSolution,
    plan_a: ScheduleSolution,
    plan_b: ScheduleSolution,
    checks: dict[str, SolutionCheck],
    certificate: dict[str, Any],
    certificate_verification: dict[str, Any],
) -> dict[str, Path]:
    outdir.mkdir(parents=True, exist_ok=True)
    schedule = schedule_table(source, interval_data, plan_a, plan_b)
    summary = summary_table(baseline, plan_a, plan_b, checks, certificate)
    certificate_summary = certificate_table(certificate, certificate_verification)
    four_hour = four_hour_table(plan_a, plan_b)
    target_purchase = target_purchase_table(source, plan_a, plan_b)

    paths = {
        "schedule_csv": outdir / "q1_schedule.csv",
        "summary_csv": outdir / "q1_summary.csv",
        "four_hour_csv": outdir / "q1_four_hour_summary.csv",
        "target_purchase_csv": outdir / "q1_target_purchase.csv",
        "workbook": outdir / "q1_results.xlsx",
        "template_result1": outdir / "result1.xlsx",
        "template_mapping_note": outdir / "result1_mapping_note.md",
    }

    schedule.to_csv(paths["schedule_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")
    summary.to_csv(paths["summary_csv"], index=False, encoding="utf-8-sig")
    four_hour.to_csv(paths["four_hour_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")
    target_purchase.to_csv(paths["target_purchase_csv"], index=False, encoding="utf-8-sig", float_format="%.10f")

    with pd.ExcelWriter(paths["workbook"], engine="xlsxwriter") as writer:
        schedule.to_excel(writer, sheet_name="schedule", index=False)
        summary.to_excel(writer, sheet_name="summary", index=False)
        four_hour.to_excel(writer, sheet_name="four_hour", index=False)
        target_purchase.to_excel(writer, sheet_name="target_purchase", index=False)
        certificate_summary.to_excel(writer, sheet_name="certificate", index=False)

        workbook = writer.book
        header_format = workbook.add_format(
            {
                "bold": True,
                "font_color": "#FFFFFF",
                "bg_color": "#234E70",
                "align": "center",
                "valign": "vcenter",
                "border": 1,
                "border_color": "#D9D9D9",
            }
        )
        number_format = workbook.add_format({"num_format": "0.000000", "valign": "vcenter"})
        text_format = workbook.add_format({"valign": "vcenter"})

        for sheet_name, frame in (
            ("schedule", schedule),
            ("summary", summary),
            ("four_hour", four_hour),
            ("target_purchase", target_purchase),
            ("certificate", certificate_summary),
        ):
            worksheet = writer.sheets[sheet_name]
            worksheet.hide_gridlines(2)
            worksheet.freeze_panes(1, 0)
            worksheet.autofilter(0, 0, len(frame), len(frame.columns) - 1)
            worksheet.set_row(0, 24, header_format)
            for column_index, column_name in enumerate(frame.columns):
                worksheet.write(0, column_index, column_name, header_format)

        schedule_ws = writer.sheets["schedule"]
        schedule_ws.set_column("A:A", 9, number_format)
        schedule_ws.set_column("B:C", 19, text_format)
        schedule_ws.set_column("D:T", 24, number_format)

        summary_ws = writer.sheets["summary"]
        summary_ws.set_column("A:A", 44, text_format)
        summary_ws.set_column("B:B", 52, text_format)

        four_hour_ws = writer.sheets["four_hour"]
        four_hour_ws.set_column("A:A", 19, text_format)
        four_hour_ws.set_column("B:G", 22, number_format)

        target_ws = writer.sheets["target_purchase"]
        target_ws.set_column("A:A", 19, text_format)
        target_ws.set_column("B:B", 10, number_format)
        target_ws.set_column("C:D", 22, number_format)

        certificate_ws = writer.sheets["certificate"]
        certificate_ws.set_column("A:A", 40, text_format)
        certificate_ws.set_column("B:B", 88, text_format)

    verify_audit_workbook(paths["workbook"], schedule, summary, certificate_summary)
    write_template_result1(paths["template_result1"], source, plan_a)
    verify_template_result1(paths["template_result1"], source, plan_a)
    write_template_mapping_note(paths["template_mapping_note"])

    return paths
